"""A napi létszámjelentés exportálása Excel és PDF formátumba.

A nehéz importok (openpyxl, reportlab) szándékosan a függvényeken belül vannak,
hogy az alkalmazás indítását ne lassítsák, és hiányuk se akadályozza a többi
funkciót — pontosan úgy, ahogy a reports modul is csinálja.
"""
from __future__ import annotations

from io import BytesIO

from .schemas import AttendanceDayRead

# A megjelenítés rögzített sorrendje; az ismeretlen státuszok a végére kerülnek.
_STATUS_ORDER = [
    "Jelen", "Szabadság", "Betegállomány", "Vezényelve",
    "Szolgálatban", "Kiküldetés", "Igazolt távollét", "Igazolatlan távollét",
]
_HEADERS = ["Név", "Rendfokozat", "Egység", "Állapot", "Megjegyzés"]


def _summary_pairs(day: AttendanceDayRead) -> list[tuple[str, int]]:
    pairs = [(s, day.summary[s]) for s in _STATUS_ORDER if day.summary.get(s)]
    extra = [(s, c) for s, c in day.summary.items() if s not in _STATUS_ORDER and c]
    return pairs + extra


def build_xlsx(day: AttendanceDayRead, unit_label: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Létszám"
    dark = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
    white_bold = Font(color="FFFFFF", bold=True)

    ws["A1"] = "Napi létszámjelentés"; ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Dátum: {day.date}"
    ws["A3"] = f"Egység: {unit_label}"
    ws["A4"] = f"Összlétszám: {day.total}"

    row = 6
    ws.cell(row=row, column=1, value="Összesítő").font = Font(bold=True); row += 1
    for status, count in _summary_pairs(day):
        ws.cell(row=row, column=1, value=status)
        ws.cell(row=row, column=2, value=count)
        row += 1
    row += 1

    for col, header in enumerate(_HEADERS, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = dark
        cell.font = white_bold
    row += 1
    for entry in day.items:
        for col, value in enumerate([entry.name, entry.rank, entry.unit, entry.status, entry.note], 1):
            ws.cell(row=row, column=col, value=value)
        row += 1

    for col, width in zip("ABCDE", [28, 18, 22, 20, 30]):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    buf.close()
    return data


def build_pdf(day: AttendanceDayRead, unit_label: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    # Arial a teljes magyar ékezetkészlethez (ő, ű); ha hiányzik, Helvetica.
    try:
        pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/Arial.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/Arialbd.ttf"))
        body_font, bold_font = "Arial", "Arial-Bold"
    except Exception:
        body_font, bold_font = "Helvetica", "Helvetica-Bold"

    C_DARK = colors.HexColor("#1e293b")
    C_WHITE = colors.white
    C_BORDER = colors.HexColor("#cbd5e1")
    C_LIGHT = colors.HexColor("#f1f5f9")
    C_MUTED = colors.HexColor("#94a3b8")

    def ps(name: str, size: int = 9, leading: int = 12, color=C_DARK, bold: bool = False) -> ParagraphStyle:
        return ParagraphStyle(name, fontName=bold_font if bold else body_font, fontSize=size, leading=leading, textColor=color)

    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=2 * cm, title="Napi létszámjelentés",
    )
    width = A4[0] - 3 * cm
    story = []

    subtitle = f"Dátum: {day.date}     Egység: {unit_label}     Összlétszám: {day.total}"
    header = Table(
        [[Paragraph("Napi létszámjelentés", ps("title", size=18, leading=22, color=C_WHITE, bold=True))],
         [Paragraph(subtitle, ps("sub", size=10, leading=14, color=C_MUTED))]],
        colWidths=[width],
    )
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), C_DARK),
        ("TOPPADDING", (0, 0), (-1, -1), 10), ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 14),
    ]))
    story.append(header)
    story.append(Spacer(1, 12))

    summary_pairs = _summary_pairs(day)
    if summary_pairs:
        rows = [[Paragraph("Összesítő", ps("sh", size=11, color=C_WHITE, bold=True)), ""]]
        rows += [[Paragraph(s, ps("c", bold=True)), Paragraph(str(c), ps("c"))] for s, c in summary_pairs]
        summary = Table(rows, colWidths=[width * 0.7, width * 0.3])
        summary.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), C_DARK),
            ("SPAN", (0, 0), (-1, 0)),
            ("GRID", (0, 1), (-1, -1), 0.5, C_BORDER),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [C_WHITE, C_LIGHT]),
            ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ("LEFTPADDING", (0, 0), (-1, -1), 8),
        ]))
        story.append(summary)
        story.append(Spacer(1, 12))

    table_rows = [[Paragraph(h, ps("th", color=C_WHITE, bold=True)) for h in _HEADERS]]
    for entry in day.items:
        table_rows.append([
            Paragraph(entry.name, ps("c")),
            Paragraph(entry.rank, ps("c")),
            Paragraph(entry.unit, ps("c")),
            Paragraph(entry.status, ps("c")),
            Paragraph(entry.note, ps("c")),
        ])
    roster = Table(table_rows, colWidths=[width * 0.26, width * 0.16, width * 0.20, width * 0.16, width * 0.22], repeatRows=1)
    roster.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), C_DARK),
        ("GRID", (0, 0), (-1, -1), 0.5, C_BORDER),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [C_WHITE, C_LIGHT]),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    story.append(roster)

    doc.build(story)
    data = buf.getvalue()
    buf.close()
    return data
