"""Leszerelési jegyzőkönyv (DOCX): mit adott vissza a katona, mi hiányzik, mennyit fizet.

A folyamat (2. betekintés, Várpalota): a leszerelő behozza a felszerelést a
laktanyába, jegyzőkönyv készül; a papír a veszprémi személyügyre és az ellátónak
megy. Ezért a dokumentumban két példány van, külön oldalon. A hiányért a katona
fizet — az összeg soronként rögzíthető. A nehéz import a függvényen belül van.
"""
from __future__ import annotations

from io import BytesIO

from pydantic import BaseModel

COPIES = ("Személyügy (Veszprém) példánya", "Ellátó példánya")


class ProtocolLine(BaseModel):
    name: str
    serialNumber: str = ""
    returned: bool = True
    amount: float = 0
    note: str = ""


def build_protocol_docx(person_name: str, rank: str, sztsz: str, unit: str, date_text: str, lines: list[ProtocolLine]) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    total = sum(line.amount for line in lines if not line.returned)

    for index, copy_name in enumerate(COPIES):
        if index:
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = title.add_run("JEGYZŐKÖNYV — leszerelés, felszerelés visszaadása")
        run.bold = True
        run.font.size = Pt(14)
        doc.add_paragraph(copy_name).alignment = WD_ALIGN_PARAGRAPH.CENTER

        for label, value in (("Név", person_name), ("Rendfokozat", rank), ("SZTSZ", sztsz), ("Egység", unit), ("Kelt", date_text)):
            row = doc.add_paragraph()
            row.add_run(f"{label}: ").bold = True
            row.add_run(value or "____________")

        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ("Felszerelés", "Gyári szám", "Visszaadva", "Hiányzik", "Fizetendő (Ft)")):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
        for line in lines:
            cells = table.add_row().cells
            cells[0].text = line.name + (f" ({line.note})" if line.note else "")
            cells[1].text = line.serialNumber
            cells[2].text = "☒" if line.returned else "☐"
            cells[3].text = "☐" if line.returned else "☒"
            cells[4].text = "" if line.returned or not line.amount else f"{line.amount:,.0f}".replace(",", " ")
        if not lines:
            table.add_row().cells[0].text = "Nincs a katonánál kiadott felszerelés."

        summary = doc.add_paragraph()
        summary.add_run(f"\nMegtérítendő összesen: {total:,.0f} Ft".replace(",", " ")).bold = True
        doc.add_paragraph("\n\n______________________          ______________________          ______________________")
        doc.add_paragraph("   leszerelő katona                      átvevő (ellátó)                      személyügy")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
