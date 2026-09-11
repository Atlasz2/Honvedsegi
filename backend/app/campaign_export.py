"""A kampányterv exportálása Excel és PDF formátumba.

Ugyanaz a felépítés, mint az attendance_export: a nehéz importok a
függvényeken belül, hogy az indítást ne lassítsák.
"""
from __future__ import annotations

from io import BytesIO

from .schemas import CampaignPlan

_HEADERS = ["Név", "Rendfokozat", "SZTSZ", "Alegység", "Állomány", "Állapot", "Jogosult", "Hiányzó követelmény"]


def _row_values(row) -> list[str]:
    return [
        row.name, row.rank, row.sztsz, row.unit, row.personStatus, row.status,
        "igen" if row.eligible else "NEM", ", ".join(row.missing),
    ]


def _subtitle(plan: CampaignPlan) -> str:
    period = f"{plan.startDate} – {plan.endDate}" if plan.endDate and plan.endDate != plan.startDate else plan.startDate
    location = f"     Helyszín: {plan.location}" if plan.location else ""
    return f"Időpont: {period}{location}"


def build_xlsx(plan: CampaignPlan) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Kampányterv"
    dark = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
    warn = PatternFill(fill_type="solid", start_color="FEF3C7", end_color="FEF3C7")
    white_bold = Font(color="FFFFFF", bold=True)

    ws["A1"] = f"Kampányterv – {plan.eventName}"; ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = _subtitle(plan)
    ws["A3"] = "Követelmények: " + (", ".join(plan.requirements) or "nincs")
    eligible = sum(1 for r in plan.rows if r.eligible)
    ws["A4"] = f"Jelentkezők: {len(plan.rows)}   Jogosult: {eligible}   Nem jogosult: {len(plan.rows) - eligible}"

    row = 6
    for col, header in enumerate(_HEADERS, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.fill = dark
        cell.font = white_bold
    row += 1
    for entry in plan.rows:
        for col, value in enumerate(_row_values(entry), 1):
            cell = ws.cell(row=row, column=col, value=value)
            if not entry.eligible:
                cell.fill = warn
        row += 1

    for col, width in zip("ABCDEFGH", [28, 16, 12, 22, 12, 14, 10, 40]):
        ws.column_dimensions[col].width = width

    buf = BytesIO()
    wb.save(buf)
    data = buf.getvalue()
    buf.close()
    return data


def build_pdf(plan: CampaignPlan) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    try:
        pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/Arial.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/Arialbd.ttf"))
        body_font, bold_font = "Arial", "Arial-Bold"
    except Exception:
        body_font, bold_font = "Helvetica", "Helvetica-Bold"

    C_DARK = colors.HexColor("#1e293b")
    C_WHITE = colors.white
    C_BORDER = colors.HexColor("#cbd5e1")
    C_LIGHT = colors.HexColor("#f1f5f9")
    C_MUTED = colors.HexColor("#94a3b8")
    C_WARN = colors.HexColor("#fef3c7")

    def ps(name: str, size: int = 8, leading: int = 10, color=C_DARK, bold: bool = False) -> ParagraphStyle:
        return ParagraphStyle(name, fontName=bold_font if bold else body_font, fontSize=size, leading=leading, textColor=color)

    page = landscape(A4)
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=page, leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=2 * cm, title=f"Kampányterv – {plan.eventName}",
    )
    width = page[0] - 3 * cm
    story = []

    eligible = sum(1 for r in plan.rows if r.eligible)
    header = Table(
        [[Paragraph(f"Kampányterv – {plan.eventName}", ps("title", size=16, leading=20, color=C_WHITE, bold=True))],
         [Paragraph(_subtitle(plan), ps("sub", size=9, leading=12, color=C_MUTED))],
         [Paragraph("Követelmények: " + (", ".join(plan.requirements) or "nincs"), ps("req", size=9, leading=12, color=C_MUTED))],
         [Paragraph(f"Jelentkezők: {len(plan.rows)}   Jogosult: {eligible}   Nem jogosult: {len(plan.rows) - eligible}",
                    ps("cnt", size=9, leading=12, color=C_WHITE))]],
        colWidths=[width],
    )
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), C_DARK),
        ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
        ("LEFTPADDING", (0, 0), (-1, -1), 14), ("RIGHTPADDING", (0, 0), (-1, -1), 14),
    ]))
    story.append(header)
    story.append(Spacer(1, 10))

    table_rows = [[Paragraph(h, ps("th", color=C_WHITE, bold=True)) for h in _HEADERS]]
    for entry in plan.rows:
        table_rows.append([Paragraph(v, ps("c")) for v in _row_values(entry)])
    fractions = [0.18, 0.10, 0.08, 0.14, 0.08, 0.09, 0.07, 0.26]
    roster = Table(table_rows, colWidths=[width * f for f in fractions], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), C_DARK),
        ("GRID", (0, 0), (-1, -1), 0.5, C_BORDER),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [C_WHITE, C_LIGHT]),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    # A nem jogosult sorok sárgák, hogy a papíron is azonnal látsszon.
    for index, entry in enumerate(plan.rows, start=1):
        if not entry.eligible:
            style.append(("BACKGROUND", (0, index), (-1, index), C_WARN))
    roster.setStyle(TableStyle(style))
    story.append(roster)

    doc.build(story)
    data = buf.getvalue()
    buf.close()
    return data
