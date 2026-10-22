"""Egyszerű táblázat-export (Excel): cím, pár fejléc-sor, fejléces tábla.

Az okmány-lista, a beléptetés és az élelmezési átjelentő közös építője. A nehéz
import a függvényen belül van (mint az attendance_export-ban).
"""
from __future__ import annotations

from io import BytesIO


def build_table_xlsx(title: str, info: list[str], headers: list[str], rows: list[list], widths: list[int] | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws["A1"] = title
    ws["A1"].font = Font(size=14, bold=True)
    for offset, line in enumerate(info, start=2):
        ws.cell(row=offset, column=1, value=line)
    header_row = len(info) + 3
    dark = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row, column=col, value=header)
        cell.fill = dark
        cell.font = Font(color="FFFFFF", bold=True)
    for r, row in enumerate(rows, start=header_row + 1):
        for col, value in enumerate(row, 1):
            ws.cell(row=r, column=col, value=value)
    for col, width in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
