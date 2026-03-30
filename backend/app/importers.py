from __future__ import annotations

import csv
import io
import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


@dataclass
class ImportRow:
    source_line: int
    data: dict[str, str]
    raw_data: dict[str, str]
    unknown_data: dict[str, str]
    enabled: bool = True


def _normalize_key(value: str) -> str:
    value = unicodedata.normalize("NFKD", value or "")
    value = "".join(ch for ch in value if not unicodedata.combining(ch))
    value = re.sub(r"[^A-Za-z0-9]+", "", value).lower()
    return value


def _to_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _decode_csv_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="ignore")


ENTITY_CONFIG: dict[str, dict[str, object]] = {
    "personnel": {
        "required": {"name", "sztsz", "rank", "unit", "status"},
        "aliases": {
            "name": "name",
            "nev": "name",
            "teljesnev": "name",
            "fullname": "name",
            "personname": "name",
            "katonanev": "name",
            "sztsz": "sztsz",
            "azonosito": "sztsz",
            "rank": "rank",
            "rendfokozat": "rank",
            "beosztas": "rank",
            "unit": "unit",
            "szervezet": "unit",
            "alegyseg": "unit",
            "status": "status",
            "statusz": "status",
            "allapot": "status",
            "allapota": "status",
            "email": "email",
            "mail": "email",
            "phone": "phone",
            "telefon": "phone",
            "birthdate": "birthDate",
            "szuletesidatum": "birthDate",
            "address": "address",
            "cim": "address",
            "joindate": "joinDate",
            "belepesdatuma": "joinDate",
            "belepes": "joinDate",
            "allomanybavetel": "joinDate",
            "notes": "notes",
            "megjegyzes": "notes",
        },
        "column_order": ["name", "sztsz", "rank", "unit", "status", "email", "phone", "birthDate", "address", "joinDate", "notes"],
    },
    "exercises": {
        "required": {"name", "type", "startDate", "endDate", "status"},
        "aliases": {
            "name": "name",
            "nev": "name",
            "gyakorlat": "name",
            "gyakorlatnev": "name",
            "gyakorlatneve": "name",
            "elnevezes": "name",
            "megnevezes": "name",
            "type": "type",
            "tipus": "type",
            "tipusa": "type",
            "gyakorlattipus": "type",
            "gyakorlattipusa": "type",
            "startdate": "startDate",
            "kezdet": "startDate",
            "kezdetdatum": "startDate",
            "kezdesdatum": "startDate",
            "kezdodatum": "startDate",
            "kezdesdatuma": "startDate",
            "start": "startDate",
            "enddate": "endDate",
            "vege": "endDate",
            "vegedatum": "endDate",
            "befejezesdatum": "endDate",
            "befejezesdatuma": "endDate",
            "end": "endDate",
            "status": "status",
            "statusz": "status",
            "allapot": "status",
            "allapota": "status",
            "location": "location",
            "helyszin": "location",
            "hely": "location",
            "terulet": "location",
            "maxpersonnel": "maxPersonnel",
            "letszam": "maxPersonnel",
            "maxletszam": "maxPersonnel",
            "maximalisletszam": "maxPersonnel",
            "resztvevokszama": "maxPersonnel",
            "description": "description",
            "leiras": "description",
            "megjegyzes": "description",
            "reszletek": "description",
        },
        "column_order": ["name", "type", "startDate", "endDate", "status", "location", "maxPersonnel", "description"],
    },
}


def _alias_similarity(source: str, target: str) -> float:
    if not source or not target:
        return 0.0
    if source == target:
        return 1.0
    if len(target) >= 4 and (source.startswith(target) or source.endswith(target) or target in source):
        return min(0.97, 0.88 + (len(target) / 100))
    if len(source) >= 4 and (target.startswith(source) or source in target):
        return min(0.93, 0.84 + (len(source) / 100))
    return SequenceMatcher(None, source, target).ratio()


def _resolve_canonical_key(entity: str, normalized_key: str) -> str | None:
    aliases = ENTITY_CONFIG[entity]["aliases"]
    if normalized_key in aliases:
        return aliases[normalized_key]

    best_score = 0.0
    second_score = 0.0
    best_canonical: str | None = None

    for alias, canonical in aliases.items():
        if len(alias) < 4:
            continue
        score = _alias_similarity(normalized_key, alias)
        if score > best_score:
            second_score = best_score
            best_score = score
            best_canonical = canonical
        elif score > second_score:
            second_score = score

    if best_score >= 0.86 and (best_score - second_score) >= 0.03:
        return best_canonical
    return None


def _map_record(entity: str, raw: dict[str, str]) -> tuple[dict[str, str], dict[str, str], dict[str, str]]:
    item: dict[str, str] = {}
    clean_raw: dict[str, str] = {}
    unknown: dict[str, str] = {}

    for key, value in raw.items():
        header = _to_text(key)
        text = _to_text(value)
        if not header:
            continue
        if text:
            clean_raw[header] = text
        canonical = _resolve_canonical_key(entity, _normalize_key(header))
        if canonical:
            item[canonical] = text
        elif text:
            unknown[header] = text

    return item, clean_raw, unknown


def _canonicalize_record(entity: str, raw: dict[str, str], require_required: bool = True) -> dict[str, str]:
    config = ENTITY_CONFIG[entity]
    required = config["required"]
    item, _, _ = _map_record(entity, raw)

    if require_required:
        missing = [field for field in required if not item.get(field)]
        if missing:
            raise ValueError(f"Hianyzik a kotelezo mezo: {', '.join(missing)}")

    return item


def _build_import_row(entity: str, source_line: int, raw: dict[str, str]) -> ImportRow | None:
    data, raw_data, unknown_data = _map_record(entity, raw)
    if not raw_data and not data and not unknown_data:
        return None
    return ImportRow(source_line=source_line, data=data, raw_data=raw_data, unknown_data=unknown_data)


def _parse_csv(entity: str, content: bytes) -> list[ImportRow]:
    text = _decode_csv_text(content)

    first_line = ""
    for line in text.splitlines():
        if line.strip():
            first_line = line
            break

    candidates = [";", ",", "\t", "|"]
    counts = {delim: first_line.count(delim) for delim in candidates}
    delimiter = max(candidates, key=lambda d: counts[d])
    if counts.get(delimiter, 0) == 0:
        sample = "\n".join(text.splitlines()[:20])
        try:
            delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
        except csv.Error:
            delimiter = ";"

    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    rows: list[ImportRow] = []
    for idx, row in enumerate(reader, start=2):
        if not row:
            continue
        mapped = {k or "": _to_text(v) for k, v in row.items()}
        parsed = _build_import_row(entity, idx, mapped)
        if parsed is None:
            continue
        rows.append(parsed)
    return rows


def _parse_xlsx(entity: str, content: bytes) -> list[ImportRow]:
    workbook = load_workbook(io.BytesIO(content), data_only=True)
    sheet = workbook.active
    header_cells = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_cells:
        return []
    headers = [_to_text(cell) for cell in header_cells]

    rows: list[ImportRow] = []
    for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        values = [_to_text(cell) for cell in row]
        if not any(values):
            continue
        data = {headers[i]: values[i] for i in range(min(len(headers), len(values)))}
        parsed = _build_import_row(entity, row_number, data)
        if parsed is None:
            continue
        rows.append(parsed)
    return rows


def _parse_table_rows(entity: str, headers: list[str], values: list[str], source_line: int) -> ImportRow | None:
    data = {headers[i]: values[i] for i in range(min(len(headers), len(values)))}
    return _build_import_row(entity, source_line, data)


def _parse_docx(entity: str, content: bytes) -> list[ImportRow]:
    doc = Document(io.BytesIO(content))
    rows: list[ImportRow] = []
    line_no = 1

    for table in doc.tables:
        if not table.rows:
            continue
        headers = [_to_text(cell.text) for cell in table.rows[0].cells]
        for r_idx, row in enumerate(table.rows[1:], start=2):
            values = [_to_text(cell.text) for cell in row.cells]
            if not any(values):
                continue
            parsed = _parse_table_rows(entity, headers, values, line_no + r_idx)
            if parsed is None:
                continue
            rows.append(parsed)
        line_no += len(table.rows) + 1

    if rows:
        return rows

    for paragraph in doc.paragraphs:
        line = _to_text(paragraph.text)
        if len(line) < 8:
            line_no += 1
            continue
        try:
            raw = _parse_line_by_pairs_raw(line)
            parsed = _build_import_row(entity, line_no, raw)
            if parsed is not None:
                rows.append(parsed)
                line_no += 1
                continue
        except Exception:
            pass

        try:
            raw = _parse_line_by_columns_raw(entity, line)
            parsed = _build_import_row(entity, line_no, raw)
            if parsed is not None:
                rows.append(parsed)
        except Exception:
            pass
        line_no += 1

    return rows


def _parse_line_by_pairs_raw(line: str) -> dict[str, str]:
    parts = re.split(r"[;|]", line)
    data: dict[str, str] = {}
    for part in parts:
        if ":" not in part:
            continue
        key, value = part.split(":", 1)
        key = key.strip()
        value = value.strip()
        if key:
            data[key] = value
    if data:
        return data
    raise ValueError("A sor nem kulcs:ertek formatu")


def _parse_line_by_columns_raw(entity: str, line: str) -> dict[str, str]:
    columns = [segment.strip() for segment in re.split(r"\t+|\s{2,}|;|\|", line) if segment.strip()]
    config = ENTITY_CONFIG[entity]
    order = config["column_order"]
    if len(columns) < len(config["required"]):
        raise ValueError("Tul keves oszlop")
    return {order[i]: columns[i] for i in range(min(len(order), len(columns)))}


def _parse_pdf(entity: str, content: bytes) -> list[ImportRow]:
    reader = PdfReader(io.BytesIO(content))
    lines: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        lines.extend(text.splitlines())

    rows: list[ImportRow] = []
    for idx, raw_line in enumerate(lines, start=1):
        line = raw_line.strip()
        if len(line) < 8:
            continue

        normalized = _normalize_key(line)
        if entity == "personnel" and "sztsz" in normalized and "nev" in normalized:
            continue
        if entity == "exercises" and "kezdet" in normalized and "vege" in normalized:
            continue

        try:
            raw = _parse_line_by_pairs_raw(line)
            parsed = _build_import_row(entity, idx, raw)
            if parsed is not None:
                rows.append(parsed)
            continue
        except Exception:
            pass

        try:
            raw = _parse_line_by_columns_raw(entity, line)
            parsed = _build_import_row(entity, idx, raw)
            if parsed is not None:
                rows.append(parsed)
        except Exception:
            continue

    return rows


def parse_import(entity: str, filename: str, content: bytes) -> list[ImportRow]:
    if entity not in ENTITY_CONFIG:
        raise ValueError("Nem tamogatott import cel")

    lower = (filename or "").lower()
    if lower.endswith(".csv"):
        return _parse_csv(entity, content)
    if lower.endswith(".xlsx") or lower.endswith(".xlsm"):
        return _parse_xlsx(entity, content)
    if lower.endswith(".pdf"):
        return _parse_pdf(entity, content)
    if lower.endswith(".docx"):
        return _parse_docx(entity, content)

    raise ValueError("Nem tamogatott formatum. Hasznalhato: .csv, .xlsx, .xlsm, .pdf, .docx")