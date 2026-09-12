"""Alapkiképzés-tábla importja: név/SZTSZ + modulonként egy oszlop dátummal.

A hadműveleti tiszt Excelje (1. betekintés): soronként egy tartalékos, a 11
modul egy-egy oszlop, a cellában a teljesítés dátuma (vagy jelölés). Itt ebből
képesítés-kiadás lesz: a modul-oszlop a BASIC_TRAINING_CATEGORY kategóriájú
képesítés-típusnak felel meg (név szerint), a cella a megszerzés dátuma. Aki
minden modult megkap, automatikusan kapja az összesítő „Alapkiképzés"-t.

Folyamat: preview (fájl → tervezet) → confirm (tervezet → adatbázis).
A tervezet a személyzet-importtal közös IMPORT_DRAFTS tárban él.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata
from datetime import date, datetime, timedelta
from typing import Any
from uuid import uuid4

from fastapi import HTTPException
from sqlalchemy import select
from sqlalchemy.orm import Session

from ..basic_training import grant_if_complete, module_types
from ..constants import BASIC_TRAINING_CATEGORY, IMPORT_DRAFT_TTL_MINUTES
from ..core.time import utc_now
from ..importers import _decode_csv_text
from ..models import PersonModel, PersonnelQualificationModel, QualificationTypeModel, new_id
from .imports import IMPORT_DRAFTS

_ENTITY = "basic_training"
_PERSON_HEADERS = {"nev": "name", "name": "name", "teljesnev": "name", "sztsz": "sztsz", "azonosito": "sztsz"}
_IGNORED_HEADERS = {"rendfokozat", "rank", "alegyseg", "szervezet", "unit", "megjegyzes", "notes", "statusz", "status"}
_DONE_MARKS = {"x", "igen", "ok", "kesz", "teljesitve", "+", "✓", "✔"}


def _norm(value: str) -> str:
    stripped = "".join(ch for ch in unicodedata.normalize("NFKD", value or "") if not unicodedata.combining(ch))
    return re.sub(r"[^a-z0-9]+", "", stripped.lower())


def _fold(value: str) -> str:
    stripped = "".join(ch for ch in unicodedata.normalize("NFD", value or "") if unicodedata.category(ch) != "Mn")
    return re.sub(r"\s+", " ", stripped).strip().lower()


def _parse_cell(value: object) -> tuple[str | None, bool]:
    """(dátum ISO vagy None, jelölve-e). Üres → nincs teljesítés."""
    if value is None:
        return None, False
    if isinstance(value, datetime):
        return value.date().isoformat(), True
    if isinstance(value, date):
        return value.isoformat(), True
    text = str(value).strip()
    if not text:
        return None, False
    if _norm(text) in _DONE_MARKS:
        return None, True
    cleaned = re.sub(r"[.\s/]+", "-", text.strip(". ")).strip("-")
    for fmt in ("%Y-%m-%d", "%Y-%m-%d-%H:%M:%S", "%Y-%m-%dT%H:%M:%S", "%d-%m-%Y"):
        try:
            return datetime.strptime(cleaned[:19], fmt).date().isoformat(), True
        except ValueError:
            continue
    # ismeretlen szöveg egy modul-cellában: teljesítésnek vesszük, dátum nélkül
    return None, True


def _read_table(filename: str, content: bytes) -> tuple[list[str], list[list[object]]]:
    lower = (filename or "").lower()
    if lower.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        ws = wb.active
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    elif lower.endswith(".csv"):
        text = _decode_csv_text(content)
        sample = text[:2048]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise HTTPException(status_code=400, detail="Támogatott formátum: .xlsx vagy .csv")
    rows = [r for r in rows if any(c not in (None, "") for c in r)]
    if not rows:
        raise HTTPException(status_code=400, detail="Üres tábla")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    return headers, rows[1:]


def _resolve_person(cells: dict[str, str], by_sztsz: dict[str, PersonModel], by_name: dict[str, list[PersonModel]]) -> tuple[PersonModel | None, str]:
    sztsz = re.sub(r"\s+", "", cells.get("sztsz", "")).upper()
    if sztsz:
        person = by_sztsz.get(sztsz)
        return (person, "") if person else (None, f"nincs ilyen SZTSZ: {sztsz}")
    name = cells.get("name", "")
    candidates = by_name.get(_fold(name), []) if name else []
    if len(candidates) == 1:
        return candidates[0], ""
    if len(candidates) > 1:
        return None, f"több {name} nevű személy — kell az SZTSZ"
    return None, f"nem található: {name or '(üres sor)'}"


def preview_basic_training(filename: str, content: bytes, db: Session) -> dict[str, Any]:
    headers, rows = _read_table(filename, content)
    person_cols: dict[int, str] = {}
    module_cols: dict[int, str] = {}
    for index, header in enumerate(headers):
        key = _norm(header)
        if not key:
            continue
        if key in _PERSON_HEADERS:
            person_cols[index] = _PERSON_HEADERS[key]
        elif key not in _IGNORED_HEADERS:
            module_cols[index] = header
    if not person_cols:
        raise HTTPException(status_code=400, detail="Nem találom a név vagy SZTSZ oszlopot a fejlécben")
    if not module_cols:
        raise HTTPException(status_code=400, detail="Nem találok modul-oszlopokat (a név/SZTSZ melletti oszlopok a modulok)")

    known = {_norm(m.name): m for m in module_types(db)}
    module_match: dict[str, str | None] = {}  # fejléc → típus-id vagy None (új)
    for header in module_cols.values():
        match = known.get(_norm(header))
        module_match[header] = match.id if match else None

    persons = db.scalars(select(PersonModel).where(PersonModel.status != "Leszerelt")).all()
    by_sztsz = {p.sztsz.upper(): p for p in persons}
    by_name: dict[str, list[PersonModel]] = {}
    for p in persons:
        by_name.setdefault(_fold(p.name), []).append(p)

    held: dict[tuple[str, str], bool] = {}
    known_ids = [m.id for m in known.values()]
    if known_ids:
        for pid, qid in db.execute(select(PersonnelQualificationModel.personnel_id, PersonnelQualificationModel.qual_type_id).where(PersonnelQualificationModel.qual_type_id.in_(known_ids))):
            held[(pid, qid)] = True

    items: list[dict[str, Any]] = []
    grants: list[dict[str, str]] = []
    unmatched: list[dict[str, Any]] = []
    already = 0
    for line_no, row in enumerate(rows, start=2):
        cells = {field: str(row[i]).strip() if i < len(row) and row[i] is not None else "" for i, field in person_cols.items()}
        person, problem = _resolve_person(cells, by_sztsz, by_name)
        completed: list[dict[str, str | None]] = []
        for index, header in module_cols.items():
            earned, marked = _parse_cell(row[index] if index < len(row) else None)
            if marked:
                completed.append({"module": header, "earnedDate": earned})
        if not person:
            unmatched.append({"line": line_no, "name": cells.get("name", ""), "sztsz": cells.get("sztsz", ""), "problem": problem, "completedCount": len(completed)})
            continue
        new_here = 0
        for entry in completed:
            qid = module_match[entry["module"]]
            if qid and held.get((person.id, qid)):
                already += 1
                continue
            new_here += 1
            grants.append({"personnelId": person.id, "module": entry["module"], "earnedDate": entry["earnedDate"] or ""})
        items.append({"line": line_no, "personnelId": person.id, "name": person.name, "sztsz": person.sztsz,
                      "completedCount": len(completed), "newCount": new_here})

    draft_id = uuid4().hex
    IMPORT_DRAFTS[draft_id] = {
        "entity": _ENTITY, "grants": grants, "module_match": module_match,
        "expires_at": utc_now() + timedelta(minutes=IMPORT_DRAFT_TTL_MINUTES),
        "rows": [], "operations": [], "created": 0, "updated": 0, "skipped": 0,
    }
    return {
        "draftId": draft_id,
        "totalRows": len(rows),
        "matchedPersons": len(items),
        "unmatched": unmatched,
        "modules": [{"header": h, "qualTypeId": qid, "known": qid is not None} for h, qid in module_match.items()],
        "unknownModules": [h for h, qid in module_match.items() if qid is None],
        "newGrants": len(grants),
        "alreadyHeld": already,
        "items": items,
    }


def confirm_basic_training(draft_id: str, create_missing_modules: bool, db: Session) -> dict[str, Any]:
    draft = IMPORT_DRAFTS.get(draft_id)
    if not draft or draft.get("entity") != _ENTITY:
        raise HTTPException(status_code=404, detail="Import draft nem található")
    if draft["expires_at"] < utc_now():
        IMPORT_DRAFTS.pop(draft_id, None)
        raise HTTPException(status_code=410, detail="Import draft lejárt")
    IMPORT_DRAFTS.pop(draft_id, None)

    module_match: dict[str, str | None] = dict(draft["module_match"])
    created_modules: list[str] = []
    for header, qid in list(module_match.items()):
        if qid is None:
            if not create_missing_modules:
                continue
            item = QualificationTypeModel(id=new_id(), name=header.strip(), category=BASIC_TRAINING_CATEGORY, validity_days=None, description="Importból létrehozott alapkiképzési modul.")
            db.add(item)
            db.flush()
            module_match[header] = item.id
            created_modules.append(header)

    today = date.today().isoformat()
    granted = 0
    skipped_unknown = 0
    touched: list[str] = []
    for grant in draft["grants"]:
        qid = module_match.get(grant["module"])
        if not qid:
            skipped_unknown += 1
            continue
        exists = db.scalar(select(PersonnelQualificationModel.id).where(
            PersonnelQualificationModel.personnel_id == grant["personnelId"], PersonnelQualificationModel.qual_type_id == qid))
        if exists:
            continue
        db.add(PersonnelQualificationModel(
            id=new_id(), personnel_id=grant["personnelId"], qual_type_id=qid,
            earned_date=grant["earnedDate"] or today, expiry_date=None, notes="Alapkiképzés-tábla import",
        ))
        granted += 1
        touched.append(grant["personnelId"])
    summaries = grant_if_complete(db, touched)
    db.commit()
    return {"granted": granted, "createdModules": created_modules, "skippedUnknownModules": skipped_unknown, "summariesGranted": summaries}
