"""A parancs dokumentumának összeállítása DOCX és PDF formátumba.

A felépítés egy általános katonai parancs-forma: fejléc (kiadó, nyilvántartási
szám), cím, tárgy, számozott fejezetek, keltezés, aláírások. A tényleges
formai követelményeket (iktatás, példányszám, minősítés) a valós minta alapján
kell véglegesíteni — ezt a C/4 nyitott kérdés hozza.

A nehéz importok a függvényeken belül vannak (mint az attendance_export-ban).
"""
from __future__ import annotations

from io import BytesIO

from .schemas import OrderRead

_SKIPPED = ("Nem szükséges",)


def _visible_chapters(order: OrderRead):
    return [ch for ch in order.chapters if ch.status not in _SKIPPED]


def _title(order: OrderRead) -> str:
    number = f"{order.number}. számú " if order.number else ""
    return f"{number}{order.typeName.upper()}"


def _dated(order: OrderRead) -> str:
    return f"Kelt: {order.issuedDate}" if order.issuedDate else "Kelt: ____________________"


def build_docx(order: OrderRead) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style.font.size = Pt(11)

    head = doc.add_paragraph()
    head.add_run(order.issuer or "").bold = True
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    meta.add_run(f"Nyt. szám: {order.number or '________'}")

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(_title(order))
    run.bold = True
    run.font.size = Pt(14)

    subject = doc.add_paragraph()
    subject.add_run("Tárgy: ").bold = True
    subject.add_run(order.subject)
    if order.personName:
        who = doc.add_paragraph()
        who.add_run("Érintett: ").bold = True
        who.add_run(order.personName)

    for index, ch in enumerate(_visible_chapters(order), start=1):
        heading = doc.add_paragraph()
        heading.add_run(f"{index}. {ch.name}").bold = True
        body = ch.content.strip() or "[A fejezet még nem készült el.]"
        for line in body.splitlines():
            doc.add_paragraph(line)

    doc.add_paragraph()
    doc.add_paragraph(_dated(order))
    doc.add_paragraph()
    for sig in order.signatures:
        block = doc.add_paragraph()
        block.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        block.add_run("______________________________\n")
        block.add_run(sig.name or "(név)").bold = True
        block.add_run(f"\n{sig.role}")

    buf = BytesIO()
    doc.save(buf)
    data = buf.getvalue()
    buf.close()
    return data


def build_pdf(order: OrderRead) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER, TA_RIGHT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    try:
        pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/Arial.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/Arialbd.ttf"))
        body_font, bold_font = "Arial", "Arial-Bold"
    except Exception:
        body_font, bold_font = "Helvetica", "Helvetica-Bold"

    def ps(name: str, size: int = 11, leading: int = 15, bold: bool = False, align=None, color=colors.black) -> ParagraphStyle:
        return ParagraphStyle(name, fontName=bold_font if bold else body_font, fontSize=size, leading=leading,
                              alignment=align if align is not None else 0, textColor=color)

    def esc(text: str) -> str:
        return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2.5 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm, title=_title(order))
    width = A4[0] - 4.5 * cm
    story = [
        Paragraph(esc(order.issuer or ""), ps("issuer", bold=True)),
        Paragraph(esc(f"Nyt. szám: {order.number or '________'}"), ps("meta", size=10, align=TA_RIGHT)),
        Spacer(1, 18),
        Paragraph(esc(_title(order)), ps("title", size=15, leading=19, bold=True, align=TA_CENTER)),
        Spacer(1, 14),
        Paragraph(f"<b>Tárgy:</b> {esc(order.subject)}", ps("subject")),
    ]
    if order.personName:
        story.append(Paragraph(f"<b>Érintett:</b> {esc(order.personName)}", ps("who")))
    story.append(Spacer(1, 10))

    for index, ch in enumerate(_visible_chapters(order), start=1):
        story.append(Paragraph(esc(f"{index}. {ch.name}"), ps("h", bold=True)))
        body = ch.content.strip() or "[A fejezet még nem készült el.]"
        story.append(Paragraph(esc(body), ps("body")))
        story.append(Spacer(1, 8))

    story.append(Spacer(1, 14))
    story.append(Paragraph(esc(_dated(order)), ps("dated")))
    story.append(Spacer(1, 24))

    if order.signatures:
        cells = []
        for sig in order.signatures:
            lines = ["______________________", f"<b>{esc(sig.name or '(név)')}</b>", esc(sig.role)]
            cells.append(Paragraph("<br/>".join(lines), ps("sig", size=10, leading=13, align=TA_CENTER)))
        table = Table([cells], colWidths=[width / len(cells)] * len(cells))
        table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(table)

    doc.build(story)
    data = buf.getvalue()
    buf.close()
    return data


def build_stats_pdf(stats: dict) -> bytes:
    """Átfutási statisztika egy lapon: összesítő, típusonként, részlegenként."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    try:
        pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/Arial.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/Arialbd.ttf"))
        body_font, bold_font = "Arial", "Arial-Bold"
    except Exception:
        body_font, bold_font = "Helvetica", "Helvetica-Bold"

    C_DARK = colors.HexColor("#1e293b")
    C_LIGHT = colors.HexColor("#f1f5f9")
    C_BORDER = colors.HexColor("#cbd5e1")
    C_MUTED = colors.HexColor("#64748b")

    def ps(name, size=10, leading=13, bold=False, align=0, color=colors.black):
        return ParagraphStyle(name, fontName=bold_font if bold else body_font, fontSize=size, leading=leading, alignment=align, textColor=color)

    def fmt(value):
        if value is None:
            return "—"
        return f"{value:.1f}" if isinstance(value, float) else str(value)

    def table(headers, rows, widths):
        data = [[Paragraph(h, ps("th", size=9, bold=True, color=colors.white)) for h in headers]]
        data += [[Paragraph(str(c), ps("c", size=9)) for c in row] for row in rows]
        t = Table(data, colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), C_DARK), ("GRID", (0, 0), (-1, -1), 0.5, C_BORDER),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, C_LIGHT]),
            ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]))
        return t

    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm, bottomMargin=2 * cm, title="Parancsok átfutása")
    width = A4[0] - 4 * cm
    totals = stats["totals"]
    period = f"{stats['year']}. év" if stats.get("year") else "teljes időszak"
    slowest = f" A leglassabb részleg: <b>{stats['slowestResponsible']}</b>." if stats.get("slowestResponsible") else ""
    type_rows = [[b["type"], b["count"], b["issued"], b["open"], b["overdue"], fmt(b["avgDays"])] for b in stats["byType"]] or [["—", 0, 0, 0, 0, "—"]]
    story = [
        Paragraph("Parancsok átfutása — statisztika", ps("title", size=16, leading=20, bold=True, align=TA_CENTER)),
        Paragraph(f"{period} · készült: {stats['generatedAt']}", ps("sub", size=9, align=TA_CENTER, color=C_MUTED)),
        Spacer(1, 14),
        Paragraph(
            f"Összesen <b>{totals['orders']}</b> parancs, ebből kiadva <b>{totals['issued']}</b>, folyamatban <b>{totals['open']}</b>, "
            f"határidőt túllépett <b>{totals['overdue']}</b>. Átlagos átfutás (létrehozástól kiadásig): <b>{fmt(totals['avgLeadDays'])} nap</b>.{slowest}",
            ps("lead", size=10, leading=14),
        ),
        Spacer(1, 12),
        Paragraph("Parancstípusonként", ps("h", size=11, bold=True)),
        Spacer(1, 4),
        table(["Típus", "Darab", "Kiadva", "Folyamatban", "Csúszott", "Átlag átfutás (nap)"], type_rows,
              [width * 0.34, width * 0.11, width * 0.11, width * 0.14, width * 0.12, width * 0.18]),
        Spacer(1, 12),
        Paragraph("Részlegenként", ps("h", size=11, bold=True)),
        Spacer(1, 4),
        table(["Részleg", "Fejezet összesen", "Kész", "Nyitott", "Lejárt határidejű", "Átlag (nap a parancs indulásától)"],
              [[r["responsible"], r["chapters"], r["done"], r["open"], r["overdue"], fmt(r["avgDays"])] for r in stats["byResponsible"]],
              [width * 0.22, width * 0.15, width * 0.11, width * 0.12, width * 0.17, width * 0.23]),
        Spacer(1, 12),
        Paragraph("Olvasat: a csúszott = a határidő után kiadott vagy határidőn túl még nyitott parancs; a részleg átlaga azt mutatja, a parancs indulásától hány nap alatt fogadta el a saját fejezetét.", ps("note", size=8, leading=11, color=C_MUTED)),
    ]
    doc.build(story)
    data = buf.getvalue()
    buf.close()
    return data
