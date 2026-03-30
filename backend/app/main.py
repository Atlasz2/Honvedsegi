from __future__ import annotations

import os
import re
from uuid import uuid4

from datetime import date, datetime, timedelta, timezone
from typing import Any, Callable, TypeVar

from fastapi import Depends, FastAPI, File, Header, HTTPException, Request, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.trustedhost import TrustedHostMiddleware
from fastapi.responses import Response
from sqlalchemy import case, func, or_, select, text
from sqlalchemy.orm import Session

from .db import Base, SessionLocal, engine, get_db
from .models import (
    ActivityLogModel,
    AnnouncementModel,
    DutyModel,
    EventModel,
    EquipmentModel,
    ExerciseModel,
    LoginAttemptModel,
    PersonModel,
    SessionTokenModel,
    SupplyModel,
    TrainingModel,
    UserModel,
    VehicleModel,
)
from .schemas import (
    ActivityLogCreate,
    ActivityLogRead,
    AnnouncementCreate,
    AnnouncementRead,
    AnnouncementUpdate,
    AuthUser,
    DutyCreate,
    DutyRead,
    DutyUpdate,
    EquipmentCheckoutRequest,
    EquipmentCreate,
    EventCreate,
    EventRead,
    EventUpdate,
    EquipmentRead,
    EquipmentUpdate,
    ExerciseCreate,
    ExerciseRead,
    ExerciseUpdate,
    LoginRequest,
    LoginResponse,
    ImportConfirmResult,
    ImportDraftUpdateRequest,
    ImportPreviewResult,
    OperationRead,
    PersonCreate,
    PersonRead,
    PersonUpdate,
    SupplyCreate,
    SupplyMovementCreate,
    SupplyRead,
    SupplyUpdate,
    TrainingCreate,
    TrainingRead,
    TrainingUpdate,
    UserCreate,
    UserRead,
    UserUpdate,
    VehicleAssignRequest,
    VehicleCreate,
    VehicleRead,
    VehicleUpdate,
)
from .security import assert_data_key_configured, assert_password_strength, decrypt_text, encrypt_text, fingerprint_token, hash_password, is_encrypted_text, issue_token, needs_rehash, verify_password
from .importers import ENTITY_CONFIG, ImportRow, parse_import
from .seed import seed_database


ModelT = TypeVar("ModelT")
SESSION_HOURS = 8
MAX_FAILED_LOGINS = 5
LOCKOUT_MINUTES = 15
GOD_USERNAME = "dev_master"
GOD_ROLE = "fejleszto"
BACKEND_ENV = os.getenv("BACKEND_ENV", "development").strip().lower()
IS_PRODUCTION = BACKEND_ENV == "production"


IMPORT_DRAFTS: dict[str, dict[str, Any]] = {}
IMPORT_DRAFT_TTL_MINUTES = 30


def _save_import_draft(
    draft_id: str,
    entity: str,
    rows: list[dict[str, Any]],
    operations: list[dict[str, Any]],
    created: int,
    updated: int,
    skipped: int,
) -> str:
    IMPORT_DRAFTS[draft_id] = {
        "entity": entity,
        "rows": rows,
        "operations": operations,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "expires_at": _utc_now() + timedelta(minutes=IMPORT_DRAFT_TTL_MINUTES),
    }
    return draft_id


def _store_import_draft(
    entity: str,
    rows: list[dict[str, Any]],
    operations: list[dict[str, Any]],
    created: int,
    updated: int,
    skipped: int,
) -> str:
    return _save_import_draft(uuid4().hex, entity, rows, operations, created, updated, skipped)


def _get_import_draft(entity: str, draft_id: str) -> dict[str, Any]:
    draft = IMPORT_DRAFTS.get(draft_id)
    if not draft:
        raise HTTPException(status_code=404, detail="Import draft nem található")
    if draft["entity"] != entity:
        raise HTTPException(status_code=400, detail="A draft más entitáshoz tartozik")
    if draft["expires_at"] < _utc_now():
        IMPORT_DRAFTS.pop(draft_id, None)
        raise HTTPException(status_code=410, detail="Import draft lejárt")
    return draft


def _pop_import_draft(entity: str, draft_id: str) -> dict[str, Any]:
    draft = _get_import_draft(entity, draft_id)
    return IMPORT_DRAFTS.pop(draft_id)


def _normalize_import_mapping(data: dict[str, Any] | None) -> dict[str, str]:
    if not data:
        return {}
    normalized: dict[str, str] = {}
    for key, value in data.items():
        text_key = str(key).strip()
        if not text_key:
            continue
        text_value = "" if value is None else str(value).strip()
        normalized[text_key] = text_value
    return normalized


def _serialize_import_row(row: ImportRow | dict[str, Any]) -> dict[str, Any]:
    if isinstance(row, ImportRow):
        return {
            "line": row.source_line,
            "enabled": row.enabled,
            "data": _normalize_import_mapping(row.data),
            "rawData": _normalize_import_mapping(row.raw_data),
            "unknownData": _normalize_import_mapping(row.unknown_data),
        }
    return {
        "line": int(row.get("line", 0)),
        "enabled": bool(row.get("enabled", True)),
        "data": _normalize_import_mapping(row.get("data")),
        "rawData": _normalize_import_mapping(row.get("rawData") or row.get("raw_data")),
        "unknownData": _normalize_import_mapping(row.get("unknownData") or row.get("unknown_data")),
    }


def _ordered_missing_fields(entity: str, data: dict[str, str]) -> list[str]:
    required = ENTITY_CONFIG[entity]["required"]
    column_order = ENTITY_CONFIG[entity]["column_order"]
    missing = [field for field in column_order if field in required and not data.get(field)]
    for field in required:
        if field not in missing and not data.get(field):
            missing.append(field)
    return missing


def _extract_validation_messages(exc: Exception) -> list[str]:
    if hasattr(exc, "errors"):
        messages: list[str] = []
        for error in exc.errors():
            loc = ".".join(str(part) for part in error.get("loc", []) if part is not None)
            label = loc or "sor"
            err_type = str(error.get("type") or "")
            if err_type == "missing":
                messages.append(f"Hianyzik a kotelezo mezo: {label}")
                continue
            if err_type == "literal_error":
                expected = str(error.get("ctx", {}).get("expected") or "")
                invalid_value = error.get("input")
                if expected:
                    messages.append(f"{label}: ervenytelen ertek ({invalid_value}). Engedelyezett: {expected}")
                else:
                    messages.append(f"{label}: ervenytelen ertek ({invalid_value})")
                continue
            messages.append(f"{label}: {error.get('msg') or 'Ervenytelen ertek'}")
        unique_messages: list[str] = []
        seen: set[str] = set()
        for message in messages:
            if message in seen:
                continue
            seen.add(message)
            unique_messages.append(message)
        return unique_messages
    text = str(exc).strip()
    return [text] if text else ["Ervenytelen sor"]


def _fallback_import_identity(entity: str, line: int, data: dict[str, str], raw_data: dict[str, str]) -> tuple[str, str]:
    if entity == "personnel":
        key = data.get("sztsz") or raw_data.get("sztsz") or raw_data.get("azonosito") or f"sor-{line}"
        name = data.get("name") or raw_data.get("nev") or raw_data.get("name") or "-"
        return key, name
    key = data.get("name") or raw_data.get("nev") or raw_data.get("name") or f"sor-{line}"
    name = data.get("name") or raw_data.get("nev") or raw_data.get("name") or "-"
    return key, name


def _evaluate_import_rows(entity: str, source_rows: list[ImportRow | dict[str, Any]], db: Session) -> dict[str, Any]:
    created = 0
    updated = 0
    skipped = 0
    issues: list[dict[str, Any]] = []
    items: list[dict[str, Any]] = []
    operations: list[dict[str, Any]] = []
    rows: list[dict[str, Any]] = []

    for source in source_rows:
        row = _serialize_import_row(source)
        rows.append(row)

        line = row["line"]
        enabled = row["enabled"]
        data = row["data"]
        raw_data = row["rawData"]
        unknown_data = row["unknownData"]
        preview_data = {k: v for k, v in data.items() if v is not None and str(v).strip()}
        preview_raw = {k: v for k, v in raw_data.items() if v is not None and str(v).strip()}
        preview_unknown = {k: v for k, v in unknown_data.items() if v is not None and str(v).strip()}
        key, name = _fallback_import_identity(entity, line, preview_data, preview_raw)
        item_issues: list[str] = []
        action = "skip"

        if not enabled:
            item_issues.append("Felhasznalo altal kihagyva")
        else:
            missing = _ordered_missing_fields(entity, preview_data)
            if missing:
                item_issues.append(f"Hianyzik a kotelezo mezo: {', '.join(missing)}")
            else:
                try:
                    if entity == "personnel":
                        payload = PersonCreate(**preview_data)
                        payload.sztsz = _normalize_sztsz(payload.sztsz)
                        existing = db.scalar(select(PersonModel).where(PersonModel.sztsz == payload.sztsz))
                        action = "update" if existing else "create"
                        key = payload.sztsz
                        name = payload.name
                    else:
                        payload = ExerciseCreate(**preview_data)
                        existing = db.scalar(
                            select(ExerciseModel).where(
                                ExerciseModel.name == payload.name,
                                ExerciseModel.start_date == payload.startDate,
                                ExerciseModel.type == payload.type,
                            )
                        )
                        action = "update" if existing else "create"
                        key = f"{payload.name}::{payload.startDate}::{payload.type}"
                        name = payload.name

                    if action == "create":
                        created += 1
                    else:
                        updated += 1

                    operations.append({
                        "entity": entity,
                        "action": action,
                        "payload": payload.model_dump(),
                    })
                except Exception as exc:
                    item_issues.extend(_extract_validation_messages(exc))

        if item_issues or not enabled:
            skipped += 1
            action = "skip"
            for message in item_issues:
                issues.append({"line": line, "message": message})

        items.append({
            "line": line,
            "action": action,
            "key": key,
            "name": name,
            "enabled": enabled,
            "data": preview_data,
            "rawData": preview_raw,
            "unknownData": preview_unknown,
            "issues": item_issues,
        })

    if not rows:
        issues.append({"line": 0, "message": "Nem sikerult ertelmezheto sort kiolvasni a fajlbol."})

    return {
        "entity": entity,
        "totalRows": len(rows),
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "issues": issues,
        "items": items,
        "operations": operations,
        "rows": rows,
    }


def _preview_response(payload: dict[str, Any], draft_id: str) -> ImportPreviewResult:
    return ImportPreviewResult(
        draftId=draft_id,
        entity=payload["entity"],
        totalRows=payload["totalRows"],
        created=payload["created"],
        updated=payload["updated"],
        skipped=payload["skipped"],
        issues=payload["issues"],
        items=payload["items"],
    )


def _required_env_csv(name: str) -> list[str]:
    raw = os.getenv(name, "").strip()
    if not raw:
        raise RuntimeError(f"Hiányzó kötelező környezeti változó production módban: {name}")
    items = [item.strip() for item in raw.split(",") if item.strip()]
    if not items:
        raise RuntimeError(f"Üres környezeti változó production módban: {name}")
    return items

app = FastAPI(
    title="Guard Guard Duty API",
    version="1.0.0",
    description="FastAPI backend katonai adminisztrációs rendszerhez SQLite adatbázissal.",
    docs_url=None if IS_PRODUCTION else "/docs",
    redoc_url=None if IS_PRODUCTION else "/redoc",
    openapi_url=None if IS_PRODUCTION else "/openapi.json",
)

if IS_PRODUCTION:
    ALLOWED_ORIGINS = _required_env_csv("BACKEND_ALLOWED_ORIGINS")
    ALLOWED_HOSTS = _required_env_csv("BACKEND_ALLOWED_HOSTS")
else:
    _allowed_origins_raw = os.getenv("BACKEND_ALLOWED_ORIGINS", "http://localhost:8080,http://127.0.0.1:8080")
    ALLOWED_ORIGINS = [origin.strip() for origin in _allowed_origins_raw.split(",") if origin.strip()]
    _allowed_hosts_raw = os.getenv("BACKEND_ALLOWED_HOSTS", "localhost,127.0.0.1")
    ALLOWED_HOSTS = [host.strip() for host in _allowed_hosts_raw.split(",") if host.strip()]

app.add_middleware(
    TrustedHostMiddleware,
    allowed_hosts=ALLOWED_HOSTS,
)

app.add_middleware(
    CORSMiddleware,
    allow_origins=ALLOWED_ORIGINS,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def security_headers_middleware(request: Request, call_next):
    response = await call_next(request)
    response.headers["X-Content-Type-Options"] = "nosniff"
    response.headers["X-Frame-Options"] = "DENY"
    response.headers["Referrer-Policy"] = "no-referrer"
    response.headers["Permissions-Policy"] = "camera=(), microphone=(), geolocation=()"
    response.headers["Cache-Control"] = "no-store"
    if request.url.scheme == "https":
        response.headers["Strict-Transport-Security"] = "max-age=63072000; includeSubDomains"
    return response


def _utc_now() -> datetime:
    return datetime.now(timezone.utc)


def _as_utc(dt: datetime) -> datetime:
    if dt.tzinfo is None:
        return dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(timezone.utc)


def _get_login_attempt(db: Session, username: str) -> LoginAttemptModel:
    attempt = db.scalar(select(LoginAttemptModel).where(LoginAttemptModel.username == username))
    if not attempt:
        attempt = LoginAttemptModel(username=username, failed_count=0)
        db.add(attempt)
        db.flush()
    return attempt


def _is_login_locked(attempt: LoginAttemptModel) -> bool:
    return bool(attempt.locked_until and _as_utc(attempt.locked_until) > _utc_now())


def _register_failed_login(db: Session, username: str) -> None:
    attempt = _get_login_attempt(db, username)
    attempt.failed_count = int(attempt.failed_count or 0) + 1
    attempt.last_failed_at = _utc_now()
    if attempt.failed_count >= MAX_FAILED_LOGINS:
        attempt.failed_count = 0
        attempt.locked_until = _utc_now() + timedelta(minutes=LOCKOUT_MINUTES)
    db.commit()


def _reset_login_attempt(db: Session, username: str) -> None:
    attempt = db.scalar(select(LoginAttemptModel).where(LoginAttemptModel.username == username))
    if not attempt:
        return
    attempt.failed_count = 0
    attempt.locked_until = None
    db.commit()


def _date_overlap(start_value: str, end_value: str, range_start: date, range_end: date) -> bool:
    start = _parse_iso_date(start_value)
    end = _parse_iso_date(end_value)
    if not start or not end:
        return False
    return start <= range_end and end >= range_start


def _parse_iso_date(value: str) -> date | None:
    if not value:
        return None
    raw = value.strip()
    if not raw:
        return None

    candidates = [raw]
    if len(raw) >= 10:
        candidates.append(raw[:10])
    if "T" in raw:
        candidates.append(raw.split("T", 1)[0])

    for candidate in candidates:
        try:
            return datetime.fromisoformat(candidate).date()
        except ValueError:
            try:
                return date.fromisoformat(candidate)
            except ValueError:
                continue
    return None


def _user_to_auth_payload(user: UserModel, expiry: datetime) -> AuthUser:
    return AuthUser(
        username=user.username,
        displayName=user.display_name,
        role=user.role,
        expiry=int(expiry.timestamp() * 1000),
    )


def _to_user_read(user: UserModel) -> UserRead:
    return UserRead(
        username=user.username,
        display_name=user.display_name,
        role=user.role,
        active=user.active,
        last_login=user.last_login,
    )



def _enc(value: str | None) -> str | None:
    if not value:
        return value
    return encrypt_text(value)


def _dec(value: str | None) -> str | None:
    if not value:
        return value
    return decrypt_text(value)

def _serialize_person(item: PersonModel) -> PersonRead:
    return PersonRead(
        id=item.id,
        name=_dec(item.name),
        sztsz=item.sztsz,
        rank=item.rank,
        unit=item.unit,
        status=item.status,
        email=_dec(item.email),
        phone=_dec(item.phone),
        birthDate=_dec(item.birth_date),
        address=_dec(item.address),
        joinDate=_dec(item.join_date),
        notes=_dec(item.notes),
    )


def _serialize_exercise(item: ExerciseModel) -> ExerciseRead:
    return ExerciseRead(
        id=item.id,
        name=item.name,
        type=item.type,
        startDate=item.start_date,
        endDate=item.end_date,
        location=item.location,
        maxPersonnel=item.max_personnel,
        description=item.description,
        status=item.status,
        assigned=item.assigned or [],
    )


def _serialize_training(item: TrainingModel) -> TrainingRead:
    return TrainingRead(
        id=item.id,
        name=item.name,
        type=item.type,
        startDate=item.start_date,
        endDate=item.end_date,
        location=item.location,
        organizer=getattr(item, "organizer", "") or "",
        maxPersonnel=item.max_personnel,
        description=item.description,
        status=item.status,
        assigned=item.assigned or [],
    )


def _serialize_event(item: EventModel) -> EventRead:
    return EventRead(
        id=item.id,
        eventType=item.event_type,
        name=item.name,
        type=item.type,
        startDate=item.start_date,
        endDate=item.end_date,
        location=item.location,
        organizer=item.organizer or "",
        maxPersonnel=item.max_personnel,
        description=item.description,
        status=item.status,
        assigned=item.assigned or [],
    )


def _serialize_equipment(item: EquipmentModel) -> EquipmentRead:
    return EquipmentRead(
        id=item.id,
        name=item.name,
        category=item.category,
        serialNumber=item.serial_number,
        qrCode=item.qr_code,
        condition=item.condition,
        description=item.description,
        checkedOutTo=item.checked_out_to,
        checkedOutToName=item.checked_out_to_name,
        checkedOutDate=item.checked_out_date,
        checkoutHistory=item.checkout_history or [],
    )


def _serialize_supply(item: SupplyModel) -> SupplyRead:
    return SupplyRead(
        id=item.id,
        name=item.name,
        category=item.category,
        unit=item.unit,
        currentQty=item.current_qty,
        minQty=item.min_qty,
        description=item.description,
        movements=item.movements or [],
    )


def _serialize_vehicle(item: VehicleModel) -> VehicleRead:
    return VehicleRead(
        id=item.id,
        plateNumber=item.plate_number,
        type=item.type,
        makeModel=item.make_model,
        year=item.year,
        km=item.km,
        nextService=item.next_service,
        nextInspection=item.next_inspection,
        status=item.status,
        notes=item.notes,
        assignedTo=item.assigned_to,
        assignedToName=item.assigned_to_name,
        serviceLog=item.service_log or [],
    )


def _serialize_duty(item: DutyModel) -> DutyRead:
    return DutyRead(
        id=item.id,
        type=item.type,
        startDate=item.start_date,
        endDate=item.end_date,
        location=item.location,
        personId=item.person_id,
        personName=item.person_name,
        notes=item.notes,
        status=item.status,
    )


def _serialize_announcement(item: AnnouncementModel) -> AnnouncementRead:
    return AnnouncementRead(
        id=item.id,
        title=item.title,
        category=item.category,
        content=item.content,
        author=item.author,
        date=item.date,
        pinned=item.pinned,
    )


def _serialize_log(item: ActivityLogModel) -> ActivityLogRead:
    return ActivityLogRead(
        id=item.id,
        timestamp=item.timestamp.isoformat(),
        userId=item.user_id,
        userName=item.user_name,
        action=item.action,
        module=item.module,
        recordName=item.record_name,
    )


def _require_model(db: Session, model_type: type[ModelT], item_id: str) -> ModelT:
    item = db.get(model_type, item_id)
    if not item:
        raise HTTPException(status_code=404, detail="Az erőforrás nem található")
    return item


def _get_current_user(
    authorization: str | None = Header(default=None),
    db: Session = Depends(get_db),
) -> UserModel:
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Bejelentkezés szükséges")

    token_value = authorization.split(" ", 1)[1].strip()
    token_key = fingerprint_token(token_value)
    session_token = db.scalar(select(SessionTokenModel).where(SessionTokenModel.token == token_key))
    if not session_token:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Érvénytelen munkamenet")
    if _as_utc(session_token.expires_at) < _utc_now():
        db.delete(session_token)
        db.commit()
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Lejárt munkamenet")

    user = db.get(UserModel, session_token.user_id)
    if not user or not user.active:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="A felhasználó nem aktív")
    return user


def _require_editor(user: UserModel = Depends(_get_current_user)) -> UserModel:
    if user.role not in {"editor", "admin", "fejleszto"}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Nincs jogosultság a művelethez")
    return user


def _require_admin(user: UserModel = Depends(_get_current_user)) -> UserModel:
    if user.role not in {"admin", "fejleszto"}:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Admin jogosultság szükséges")
    return user


def _is_god_user(user: UserModel) -> bool:
    return user.username == GOD_USERNAME and user.role == GOD_ROLE and user.active


def _require_god_user(user: UserModel = Depends(_get_current_user)) -> UserModel:
    if not _is_god_user(user):
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Csak a dev_master jogosult erre a műveletre")
    return user


def _enforce_single_god_user(db: Session) -> None:
    god_user = db.scalar(select(UserModel).where(UserModel.username == GOD_USERNAME))
    if not god_user:
        dev_pwd = os.getenv("BACKEND_DEV_MASTER_PASSWORD", "").strip()
        if not dev_pwd:
            raise RuntimeError("Hiányzó BACKEND_DEV_MASTER_PASSWORD a dev_master létrehozásához")
        god_user = UserModel(
            username=GOD_USERNAME,
            password_hash=hash_password(dev_pwd),
            display_name="Fejlesztő Mester",
            role=GOD_ROLE,
            active=True,
            protected=True,
        )
        db.add(god_user)

    god_user.role = GOD_ROLE
    god_user.active = True
    god_user.protected = True

    other_devs = db.scalars(
        select(UserModel).where(UserModel.role == GOD_ROLE, UserModel.username != GOD_USERNAME)
    ).all()
    for user in other_devs:
        user.role = "admin"
        user.protected = False

    db.commit()


def _normalize_sztsz(value: str) -> str:
    normalized = re.sub(r"\s+", "", (value or "").strip()).upper()
    if re.fullmatch(r"\d{8}", normalized):
        return normalized
    if re.fullmatch(r"[A-Z]{2}\d{6}", normalized):
        return normalized
    raise HTTPException(status_code=400, detail="Az SZTSz formátuma 8 számjegy vagy 2 betű + 6 számjegy lehet")


def _assert_unique_sztsz(db: Session, sztsz: str, exclude_id: str | None = None) -> None:
    query = select(PersonModel).where(PersonModel.sztsz == sztsz)
    existing = db.scalar(query)
    if existing and existing.id != exclude_id:
        raise HTTPException(status_code=409, detail="Ez az SZTSz már létezik")


def _ensure_personnel_sztsz_schema(db: Session) -> None:
    columns = {row[1] for row in db.execute(text("PRAGMA table_info(personnel)")).fetchall()}
    if "sztsz" not in columns:
        db.execute(text("ALTER TABLE personnel ADD COLUMN sztsz TEXT"))

    rows = db.execute(text("SELECT id, sztsz FROM personnel ORDER BY id")).fetchall()
    used: set[str] = set()
    next_value = 10000000

    for person_id, sztsz in rows:
        normalized = str(sztsz).strip() if sztsz is not None else ""
        valid = len(normalized) == 8 and normalized.isdigit() and normalized not in used
        if valid:
            used.add(normalized)
            continue

        while True:
            candidate = f"{next_value:08d}"
            next_value += 1
            if candidate not in used:
                break

        used.add(candidate)
        db.execute(text("UPDATE personnel SET sztsz = :sztsz WHERE id = :id"), {"sztsz": candidate, "id": person_id})

    db.execute(text("CREATE UNIQUE INDEX IF NOT EXISTS ix_personnel_sztsz ON personnel(sztsz)"))
    db.commit()


def _apply_person(target: PersonModel, payload: PersonCreate | PersonUpdate) -> None:
    target.name = encrypt_text(payload.name)
    target.sztsz = payload.sztsz
    target.rank = payload.rank
    target.unit = payload.unit
    target.status = payload.status
    target.email = _enc(payload.email)
    target.phone = _enc(payload.phone)
    target.birth_date = _enc(payload.birthDate)
    target.address = _enc(payload.address)
    target.join_date = _enc(payload.joinDate)
    target.notes = _enc(payload.notes)


def _apply_exercise(target: ExerciseModel, payload: ExerciseCreate | ExerciseUpdate) -> None:
    target.name = payload.name
    target.type = payload.type
    target.start_date = payload.startDate
    target.end_date = payload.endDate
    target.location = payload.location
    target.max_personnel = payload.maxPersonnel
    target.description = payload.description
    target.status = payload.status
    target.assigned = [item.model_dump() for item in payload.assigned]


def _apply_training(target: TrainingModel, payload: TrainingCreate | TrainingUpdate) -> None:
    target.name = payload.name
    target.type = payload.type
    target.start_date = payload.startDate
    target.end_date = payload.endDate
    target.location = payload.location
    if hasattr(target, "organizer"):
        target.organizer = payload.organizer
    target.max_personnel = payload.maxPersonnel
    target.description = payload.description
    target.status = payload.status
    target.assigned = [item.model_dump() for item in payload.assigned]


def _apply_event(target: EventModel, payload: EventCreate | EventUpdate) -> None:
    target.event_type = payload.eventType
    target.name = payload.name
    target.type = payload.type
    target.start_date = payload.startDate
    target.end_date = payload.endDate
    target.location = payload.location
    target.organizer = payload.organizer
    target.max_personnel = payload.maxPersonnel
    target.description = payload.description
    target.status = payload.status
    target.assigned = payload.assigned


def _apply_equipment(target: EquipmentModel, payload: EquipmentCreate | EquipmentUpdate) -> None:
    target.name = payload.name
    target.category = payload.category
    target.serial_number = payload.serialNumber
    target.qr_code = payload.qrCode
    target.condition = payload.condition
    target.description = payload.description
    target.checked_out_to = payload.checkedOutTo
    target.checked_out_to_name = payload.checkedOutToName
    target.checked_out_date = payload.checkedOutDate
    target.checkout_history = [item.model_dump() for item in payload.checkoutHistory]


def _apply_supply(target: SupplyModel, payload: SupplyCreate | SupplyUpdate) -> None:
    target.name = payload.name
    target.category = payload.category
    target.unit = payload.unit
    target.current_qty = payload.currentQty
    target.min_qty = payload.minQty
    target.description = payload.description
    target.movements = [item.model_dump() for item in payload.movements]


def _apply_vehicle(target: VehicleModel, payload: VehicleCreate | VehicleUpdate) -> None:
    target.plate_number = payload.plateNumber
    target.type = payload.type
    target.make_model = payload.makeModel
    target.year = payload.year
    target.km = payload.km
    target.next_service = payload.nextService
    target.next_inspection = payload.nextInspection
    target.status = payload.status
    target.notes = payload.notes
    target.assigned_to = payload.assignedTo
    target.assigned_to_name = payload.assignedToName
    target.service_log = [item.model_dump() for item in payload.serviceLog]


def _apply_duty(target: DutyModel, payload: DutyCreate | DutyUpdate) -> None:
    target.type = payload.type
    target.start_date = payload.startDate
    target.end_date = payload.endDate
    target.location = payload.location
    target.person_id = payload.personId
    target.person_name = payload.personName
    target.notes = payload.notes
    target.status = payload.status



_PERSON_ENC_FIELDS = ("name", "email", "phone", "birth_date", "address", "join_date", "notes")


def _ensure_personnel_encryption(db: Session) -> None:
    if not os.getenv("BACKEND_DATA_KEY", "").strip():
        return
    rows = db.scalars(select(PersonModel)).all()
    changed = False
    for person in rows:
        for field in _PERSON_ENC_FIELDS:
            value = getattr(person, field)
            if value and not is_encrypted_text(value):
                setattr(person, field, encrypt_text(value))
                changed = True
    if changed:
        db.commit()

@app.on_event("startup")
def on_startup() -> None:
    assert_data_key_configured(IS_PRODUCTION)
    Base.metadata.create_all(bind=engine)
    with SessionLocal() as db:
        seed_database(db)
        _ensure_personnel_sztsz_schema(db)
        _ensure_personnel_encryption(db)
        _enforce_single_god_user(db)


@app.get("/api/health")
def health(db: Session = Depends(get_db)) -> dict[str, str]:
    try:
        db.execute(text("SELECT 1"))
    except Exception as exc:
        raise HTTPException(status_code=503, detail=f"Adatbázis nem elérhető: {exc}") from exc

    return {
        "status": "ok",
        "environment": BACKEND_ENV,
        "time": _utc_now().isoformat(),
    }

@app.get("/api/operations/summary")
def operations_summary(
    base_date: str | None = None,
    db: Session = Depends(get_db),
    _: UserModel = Depends(_get_current_user),
):
    if base_date:
        parsed_base = _parse_iso_date(base_date)
        if not parsed_base:
            raise HTTPException(status_code=400, detail="Érvénytelen base_date formátum")
        base = parsed_base
    else:
        base = _utc_now().date()

    next_week_end = base + timedelta(days=7)
    plus14_day = base + timedelta(days=14)

    exercise_items = db.scalars(
        select(ExerciseModel).where(ExerciseModel.status.in_(["Tervezett", "Folyamatban"]))
    ).all()
    duty_items = db.scalars(
        select(DutyModel).where(DutyModel.status.in_(["Tervezett", "Teljesített"]))
    ).all()

    shooting_keywords = ["lőtér", "loter", "loter"]
    next_week_shooting = []
    for item in exercise_items:
        start = _parse_iso_date(item.start_date)
        end = _parse_iso_date(item.end_date)
        if not start or not end:
            continue
        if end < base or start > next_week_end:
            continue
        location_lower = (item.location or "").lower()
        if not any(keyword in location_lower for keyword in shooting_keywords):
            continue
        next_week_shooting.append(
            {
                "id": item.id,
                "name": item.name,
                "startDate": item.start_date,
                "endDate": item.end_date,
                "location": item.location,
                "status": item.status,
                "assignedCount": len(item.assigned or []),
                "maxPersonnel": item.max_personnel,
            }
        )

    plus14_duties = []
    for item in duty_items:
        start = _parse_iso_date(item.start_date)
        end = _parse_iso_date(item.end_date)
        if not start or not end:
            continue
        if start <= plus14_day <= end:
            plus14_duties.append(
                {
                    "id": item.id,
                    "type": item.type,
                    "startDate": item.start_date,
                    "endDate": item.end_date,
                    "location": item.location,
                    "personId": item.person_id,
                    "personName": item.person_name,
                    "status": item.status,
                }
            )

    return {
        "baseDate": base.isoformat(),
        "nextWeekEnd": next_week_end.isoformat(),
        "plus14Date": plus14_day.isoformat(),
        "nextWeekShooting": {
            "count": len(next_week_shooting),
            "items": sorted(next_week_shooting, key=lambda x: x["startDate"]),
        },
        "plus14Duties": {
            "count": len(plus14_duties),
            "items": sorted(plus14_duties, key=lambda x: x["startDate"]),
        },
    }


def _report_title(template: str) -> str:
    title_map = {
        "overview": "Összesített műveleti riport",
        "operations": "Műveleti naptár riport",
        "duties": "Szolgálati kivonat",
        "events": "Eseménynaptár riport",
        "focus": "Részletes fókusz riport",
    }
    return title_map[template]


def _serialize_report_list_item(item: Any, item_type: str) -> dict[str, Any]:
    if item_type == "duty":
        return {
            "id": item.id,
            "itemType": item_type,
            "type": item.type,
            "personId": item.person_id,
            "personName": item.person_name,
            "startDate": item.start_date,
            "endDate": item.end_date,
            "location": item.location or "-",
            "status": item.status,
            "previewRow": f"- {item.start_date} -> {item.end_date} | {item.type} | {item.person_name} | {item.location or '-'} | {item.status}",
        }

    assigned = getattr(item, "assigned", None) or []
    payload = {
        "id": item.id,
        "itemType": item_type,
        "name": item.name,
        "type": getattr(item, "type", "-") or "-",
        "startDate": item.start_date,
        "endDate": item.end_date,
        "location": item.location or "-",
        "status": item.status,
        "maxPersonnel": getattr(item, "max_personnel", 0),
        "assignedCount": len(assigned),
        "previewRow": f"- {item.start_date} -> {item.end_date} | {item.name} | {item.location or '-'} | {item.status}",
    }
    if hasattr(item, "organizer"):
        payload["organizer"] = getattr(item, "organizer", "") or "-"
    return payload


def _serialize_focus_report(item: Any, focus_type: str, focus_id: str) -> dict[str, Any]:
    if focus_type == "duty":
        return {
            "type": focus_type,
            "id": focus_id,
            "headline": f"Szolgálat: {item.type}",
            "description": "",
            "participants": [],
            "details": [
                {"label": "Időszak", "value": f"{item.start_date} - {item.end_date}"},
                {"label": "Személy", "value": item.person_name},
                {"label": "Helyszín", "value": item.location or "-"},
                {"label": "Státusz", "value": item.status},
            ],
        }

    assigned = getattr(item, "assigned", None) or []
    participants = []
    for entry in assigned[:250]:
        participants.append(
            {
                "personName": entry.get("personName") or entry.get("person_name") or "-",
                "detail": entry.get("role") or entry.get("attendance") or "-",
            }
        )

    details = [
        {"label": "Típus", "value": getattr(item, "type", "-") or "-"},
        {"label": "Időszak", "value": f"{item.start_date} - {item.end_date}"},
        {"label": "Helyszín", "value": item.location or "-"},
        {"label": "Státusz", "value": item.status},
        {"label": "Max. létszám", "value": str(getattr(item, "max_personnel", 0))},
        {"label": "Hozzárendelve", "value": f"{len(assigned)} fő"},
    ]
    if hasattr(item, "organizer"):
        details.insert(4, {"label": "Szervező", "value": getattr(item, "organizer", "") or "-"})

    return {
        "type": focus_type,
        "id": focus_id,
        "headline": f"Megnevezés: {item.name}",
        "description": getattr(item, "description", "") or "",
        "participants": participants,
        "details": details,
    }


def _build_operations_report_data(
    db: Session,
    date_from: str | None,
    date_to: str | None,
    template: str,
    focus_type: str | None,
    focus_id: str | None,
) -> dict[str, Any]:
    allowed_templates = {"overview", "operations", "duties", "events", "focus"}
    allowed_focus_types = {"exercise", "training", "event", "duty"}
    if template not in allowed_templates:
        raise HTTPException(status_code=400, detail="Nem támogatott riportminta")
    if focus_type and focus_type not in allowed_focus_types:
        raise HTTPException(status_code=400, detail="Nem támogatott fókusz típus")

    start_date = _parse_iso_date(date_from) if date_from else _utc_now().date()
    end_date = _parse_iso_date(date_to) if date_to else start_date + timedelta(days=30)

    if not start_date or not end_date:
        raise HTTPException(status_code=400, detail="Érvénytelen dátumtartomány")
    if end_date < start_date:
        raise HTTPException(status_code=400, detail="A záró dátum nem lehet korábbi")

    exercise_items = [
        item for item in db.scalars(select(ExerciseModel).order_by(ExerciseModel.start_date)).all()
        if _date_overlap(item.start_date, item.end_date, start_date, end_date)
    ]
    training_items = [
        item for item in db.scalars(select(TrainingModel).order_by(TrainingModel.start_date)).all()
        if _date_overlap(item.start_date, item.end_date, start_date, end_date)
    ]
    event_items = [
        item for item in db.scalars(select(EventModel).order_by(EventModel.start_date)).all()
        if _date_overlap(item.start_date, item.end_date, start_date, end_date)
    ]
    duty_items = [
        item for item in db.scalars(select(DutyModel).order_by(DutyModel.start_date)).all()
        if _date_overlap(item.start_date, item.end_date, start_date, end_date)
    ]

    sections: list[dict[str, Any]] = []

    def add_section(key: str, title: str, items: list[Any], item_type: str, limit: int):
        visible_items = items[:limit]
        sections.append(
            {
                "key": key,
                "title": title,
                "count": len(items),
                "truncated": len(items) > len(visible_items),
                "items": [_serialize_report_list_item(item, item_type) for item in visible_items],
            }
        )

    focus_payload = None
    if template == "focus":
        if not focus_type or not focus_id:
            raise HTTPException(status_code=400, detail="A fókusz riporthoz típus és azonosító szükséges")

        item = None
        if focus_type == "exercise":
            item = db.scalar(select(ExerciseModel).where(ExerciseModel.id == focus_id))
        elif focus_type == "training":
            item = db.scalar(select(TrainingModel).where(TrainingModel.id == focus_id))
        elif focus_type == "event":
            item = db.scalar(select(EventModel).where(EventModel.id == focus_id))
        elif focus_type == "duty":
            item = db.scalar(select(DutyModel).where(DutyModel.id == focus_id))

        if not item:
            raise HTTPException(status_code=404, detail="A kiválasztott rekord nem található")

        focus_payload = _serialize_focus_report(item, focus_type, focus_id)
    else:
        if template in {"overview", "operations"}:
            add_section("exercises", "Gyakorlatok", exercise_items, "exercise", 300)
            add_section("trainings", "Kiképzések", training_items, "training", 300)
        if template == "overview":
            add_section("events", "Események", event_items, "event", 300)
            add_section("duties", "Szolgálatok", duty_items, "duty", 400)
        elif template == "duties":
            add_section("duties", "Szolgálatok", duty_items, "duty", 400)
        elif template == "events":
            add_section("events", "Események", event_items, "event", 300)

    return {
        "template": template,
        "title": _report_title(template),
        "interval": {
            "dateFrom": start_date.isoformat(),
            "dateTo": end_date.isoformat(),
        },
        "focusType": focus_type,
        "focusId": focus_id,
        "summary": {
            "exercises": len(exercise_items),
            "trainings": len(training_items),
            "events": len(event_items),
            "duties": len(duty_items),
        },
        "sections": sections,
        "focus": focus_payload,
    }


@app.get("/api/reports/operations/preview")
def operations_report_preview(
    date_from: str | None = None,
    date_to: str | None = None,
    template: str = "overview",
    focus_type: str | None = None,
    focus_id: str | None = None,
    db: Session = Depends(get_db),
    _: UserModel = Depends(_get_current_user),
):
    return _build_operations_report_data(db, date_from, date_to, template, focus_type, focus_id)


@app.get("/api/reports/operations.pdf")
def operations_pdf_report(
    date_from: str | None = None,
    date_to: str | None = None,
    template: str = "overview",
    focus_type: str | None = None,
    focus_id: str | None = None,
    db: Session = Depends(get_db),
    _: UserModel = Depends(_get_current_user),
):
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib import colors
        from reportlab.lib.units import cm
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.platypus import (
            SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
            HRFlowable, KeepTogether,
        )
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"PDF modul hiba: {exc}") from exc

    # ── Fontok regisztrálása (magyar karakterek) ──────────────────────────────
    try:
        pdfmetrics.registerFont(TTFont("Arial", "C:/Windows/Fonts/Arial.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Bold", "C:/Windows/Fonts/Arialbd.ttf"))
        pdfmetrics.registerFont(TTFont("Arial-Italic", "C:/Windows/Fonts/Ariali.ttf"))
        pdfmetrics.registerFontFamily("Arial", normal="Arial", bold="Arial-Bold", italic="Arial-Italic")
        body_font = "Arial"
        bold_font = "Arial-Bold"
    except Exception:
        body_font = "Helvetica"
        bold_font = "Helvetica-Bold"

    # ── Szín paletta ──────────────────────────────────────────────────────────
    C_DARK   = colors.HexColor("#1e293b")   # fejléc háttér
    C_MED    = colors.HexColor("#334155")   # szekció fejléc háttér
    C_LIGHT  = colors.HexColor("#f1f5f9")   # sáv-háttér
    C_ACCENT = colors.HexColor("#3b82f6")   # kiemelő kék
    C_WHITE  = colors.white
    C_BORDER = colors.HexColor("#cbd5e1")

    # ── Stílusok ──────────────────────────────────────────────────────────────
    def ps(name, font=body_font, size=9, leading=12, color=C_DARK, bold=False, **kw):
        return ParagraphStyle(
            name, fontName=bold_font if bold else font,
            fontSize=size, leading=leading, textColor=color, **kw
        )

    sTitle    = ps("title",   size=18, leading=22, bold=True,  color=C_WHITE)
    sSub      = ps("sub",     size=10, leading=14, color=colors.HexColor("#94a3b8"))
    sSection  = ps("sec",     size=11, leading=14, bold=True,  color=C_WHITE)
    sLabel    = ps("label",   size=8,  leading=11, bold=True,  color=C_MED)
    sValue    = ps("value",   size=9,  leading=12, color=C_DARK)
    sCell     = ps("cell",    size=8,  leading=11, color=C_DARK)
    sCellBold = ps("cellb",   size=8,  leading=11, bold=True,  color=C_DARK)
    sEmpty    = ps("empty",   size=8,  leading=11, color=colors.HexColor("#94a3b8"))
    sDesc     = ps("desc",    size=8,  leading=12, color=C_DARK)

    from io import BytesIO
    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        leftMargin=1.5*cm, rightMargin=1.5*cm,
        topMargin=1.5*cm,  bottomMargin=2*cm,
        title="Guard Guard Duty – Riport",
    )
    W = A4[0] - 3*cm   # hasznos szélesség

    story = []
    report_data = _build_operations_report_data(db, date_from, date_to, template, focus_type, focus_id)
    interval    = report_data["interval"]
    start_date  = interval["dateFrom"]
    end_date    = interval["dateTo"]
    summary     = report_data["summary"]

    # ── Fejléc banner ─────────────────────────────────────────────────────────
    subtitle_text = f"Intervallum: {start_date}  –  {end_date}"
    if template == "focus":
        subtitle_text += f"   |   Fókusz: {focus_type or '–'}"

    header_table = Table(
        [[Paragraph(report_data["title"], sTitle)],
         [Paragraph(subtitle_text, sSub)]],
        colWidths=[W],
    )
    header_table.setStyle(TableStyle([
        ("BACKGROUND", (0,0), (-1,-1), C_DARK),
        ("ROWBACKGROUNDS", (0,0), (-1,-1), [C_DARK, C_DARK]),
        ("TOPPADDING",    (0,0), (-1,-1), 10),
        ("BOTTOMPADDING", (0,0), (-1,-1), 8),
        ("LEFTPADDING",   (0,0), (-1,-1), 14),
        ("RIGHTPADDING",  (0,0), (-1,-1), 14),
        ("ROUNDEDCORNERS", [6]),
    ]))
    story.append(header_table)
    story.append(Spacer(1, 0.4*cm))

    # ── Összesítő kártyák ─────────────────────────────────────────────────────
    summary_labels = [
        ("Gyakorlatok",  str(summary["exercises"])),
        ("Kiképzések",   str(summary["trainings"])),
        ("Események",    str(summary["events"])),
        ("Szolgálatok",  str(summary["duties"])),
    ]
    card_w = W / 4 - 0.1*cm
    card_data = [[
        Table([[Paragraph(v, ps("cv", size=20, leading=24, bold=True, color=C_ACCENT))],
               [Paragraph(l, ps("cl", size=8,  leading=10, color=colors.HexColor("#64748b")))]],
              colWidths=[card_w])
        for l, v in summary_labels
    ]]
    card_styles = []
    for col in range(4):
        card_styles += [
            ("BACKGROUND",    (col,0), (col,0), C_LIGHT),
            ("BOX",           (col,0), (col,0), 0.5, C_BORDER),
            ("TOPPADDING",    (col,0), (col,0), 8),
            ("BOTTOMPADDING", (col,0), (col,0), 8),
            ("LEFTPADDING",   (col,0), (col,0), 10),
            ("RIGHTPADDING",  (col,0), (col,0), 10),
        ]
    cards_table = Table(card_data, colWidths=[card_w + 0.1*cm]*4)
    cards_table.setStyle(TableStyle(card_styles))
    story.append(cards_table)
    story.append(Spacer(1, 0.5*cm))

    # ── Segédfüggvények ───────────────────────────────────────────────────────
    def section_header(title: str, count: int | None = None):
        label = title if count is None else f"{title}  ({count} db)"
        t = Table([[Paragraph(label, sSection)]], colWidths=[W])
        t.setStyle(TableStyle([
            ("BACKGROUND",    (0,0), (-1,-1), C_MED),
            ("TOPPADDING",    (0,0), (-1,-1), 7),
            ("BOTTOMPADDING", (0,0), (-1,-1), 7),
            ("LEFTPADDING",   (0,0), (-1,-1), 12),
            ("RIGHTPADDING",  (0,0), (-1,-1), 12),
        ]))
        return t

    def info_row(label: str, value: str):
        return Table(
            [[Paragraph(label, sLabel), Paragraph(value, sValue)]],
            colWidths=[3.5*cm, W - 3.5*cm],
        )

    def data_table(headers: list[str], rows: list[list[str]], col_widths: list[float]):
        h_cells = [Paragraph(h, sCellBold) for h in headers]
        body    = [h_cells] + [[Paragraph(str(c), sCell) for c in row] for row in rows]
        styles  = [
            ("BACKGROUND",    (0,0), (-1,0),  C_DARK),
            ("TEXTCOLOR",     (0,0), (-1,0),  C_WHITE),
            ("FONTNAME",      (0,0), (-1,0),  bold_font),
            ("FONTSIZE",      (0,0), (-1,0),  8),
            ("ROWBACKGROUNDS",(0,1), (-1,-1), [C_WHITE, C_LIGHT]),
            ("GRID",          (0,0), (-1,-1), 0.3, C_BORDER),
            ("TOPPADDING",    (0,0), (-1,-1), 4),
            ("BOTTOMPADDING", (0,0), (-1,-1), 4),
            ("LEFTPADDING",   (0,0), (-1,-1), 6),
            ("RIGHTPADDING",  (0,0), (-1,-1), 6),
            ("VALIGN",        (0,0), (-1,-1), "MIDDLE"),
        ]
        t = Table(body, colWidths=col_widths, repeatRows=1)
        t.setStyle(TableStyle(styles))
        return t

    # ── Tartalom: fókusz riport ───────────────────────────────────────────────
    if template == "focus":
        focus = report_data["focus"]
        story.append(KeepTogether([
            section_header(focus["headline"]),
            Spacer(1, 0.2*cm),
        ]))
        for det in focus["details"]:
            story.append(info_row(det["label"], det["value"]))
            story.append(Spacer(1, 0.15*cm))
        if focus["description"]:
            story.append(Spacer(1, 0.3*cm))
            story.append(section_header("Leírás"))
            story.append(Spacer(1, 0.2*cm))
            for raw_line in focus["description"].splitlines():
                story.append(Paragraph(raw_line or " ", sDesc))
        if focus["participants"]:
            story.append(Spacer(1, 0.4*cm))
            story.append(section_header("Résztvevők", len(focus["participants"])))
            story.append(Spacer(1, 0.2*cm))
            rows = [[e["personName"], e["detail"]] for e in focus["participants"]]
            story.append(data_table(
                ["Név", "Részleg / szerep"],
                rows,
                [W * 0.55, W * 0.45],
            ))

    # ── Tartalom: lista riportok ──────────────────────────────────────────────
    else:
        col_cfg = {
            "duty":     (["Kezdés", "Vége", "Típus", "Személy", "Helyszín", "Státusz"],
                         [2.2*cm, 2.2*cm, 2.8*cm, 4.5*cm, 3.5*cm, 2.5*cm]),
            "exercise": (["Kezdés", "Vége", "Megnevezés", "Helyszín", "Státusz"],
                         [2.2*cm, 2.2*cm, 5.5*cm, 4.0*cm, 3.5*cm]),
            "training": (["Kezdés", "Vége", "Megnevezés", "Helyszín", "Státusz"],
                         [2.2*cm, 2.2*cm, 5.5*cm, 4.0*cm, 3.5*cm]),
            "event":    (["Kezdés", "Vége", "Megnevezés", "Helyszín", "Státusz"],
                         [2.2*cm, 2.2*cm, 5.5*cm, 4.0*cm, 3.5*cm]),
        }
        type_map = {
            "duty":     "duty",
            "exercises":"exercise",
            "trainings":"training",
            "events":   "event",
            "duties":   "duty",
        }
        for sec in report_data["sections"]:
            item_type = type_map.get(sec["key"], "exercise")
            headers, widths = col_cfg.get(item_type, col_cfg["exercise"])

            story.append(Spacer(1, 0.3*cm))
            story.append(section_header(sec["title"], sec["count"]))
            story.append(Spacer(1, 0.2*cm))

            if not sec["items"]:
                story.append(Paragraph("Nincs találat a megadott feltételekre.", sEmpty))
                story.append(Spacer(1, 0.2*cm))
                continue

            def _row(item: dict, itype: str) -> list[str]:
                if itype == "duty":
                    return [
                        item.get("startDate",""), item.get("endDate",""),
                        item.get("type",""), item.get("personName",""),
                        item.get("location",""), item.get("status",""),
                    ]
                return [
                    item.get("startDate",""), item.get("endDate",""),
                    item.get("name",""), item.get("location",""), item.get("status",""),
                ]

            rows = [_row(it, item_type) for it in sec["items"]]
            story.append(data_table(headers, rows, widths))
            if sec.get("truncated"):
                story.append(Paragraph(
                    f"(Csak az első {len(sec['items'])} sor látható – szűkítsd az intervallumot a teljes listához.)",
                    sEmpty,
                ))
            story.append(Spacer(1, 0.1*cm))

    # ── Lábléc (oldalszám egyszerű megoldással) ───────────────────────────────
    from reportlab.platypus import PageBreak
    def _add_page_number(canvas_obj, doc_obj):
        canvas_obj.saveState()
        canvas_obj.setFont(body_font, 7)
        canvas_obj.setFillColor(colors.HexColor("#94a3b8"))
        page_num = f"{canvas_obj.getPageNumber()}. oldal  |  Guard Guard Duty – {report_data['title']}"
        canvas_obj.drawRightString(A4[0] - 1.5*cm, 1.2*cm, page_num)
        canvas_obj.restoreState()

    doc.build(story, onFirstPage=_add_page_number, onLaterPages=_add_page_number)
    data = buffer.getvalue()
    buffer.close()

    filename_map = {
        "overview": "osszesitett-muveleti-riport.pdf",
        "operations": "muveleti-naptar-riport.pdf",
        "duties": "szolgalati-kivonat.pdf",
        "events": "esemenynaptar-riport.pdf",
        "focus": f"fokusz-riport-{focus_type or 'elem'}.pdf",
    }
    return Response(
        content=data,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename_map.get(template, 'riport.pdf')}"},
    )
@app.post("/api/auth/login", response_model=LoginResponse)
def login(payload: LoginRequest, db: Session = Depends(get_db)) -> LoginResponse:
    normalized_username = (payload.username or "").strip()
    if not normalized_username:
        raise HTTPException(status_code=401, detail="Hibás felhasználónév vagy jelszó")

    attempt = _get_login_attempt(db, normalized_username)
    if _is_login_locked(attempt):
        raise HTTPException(status_code=429, detail="Túl sok hibás próbálkozás. Próbáld újra később.")

    user = db.scalar(select(UserModel).where(UserModel.username == normalized_username))
    if not user or not verify_password(payload.password, user.password_hash):
        _register_failed_login(db, normalized_username)
        raise HTTPException(status_code=401, detail="Hibás felhasználónév vagy jelszó")
    if not user.active:
        raise HTTPException(status_code=403, detail="A felhasználó inaktív")

    if needs_rehash(user.password_hash):
        try:
            assert_password_strength(payload.password)
            user.password_hash = hash_password(payload.password)
        except ValueError:
            # Régi, gyenge jelszó esetén ne bukjon el a login; a rehash ilyenkor elmarad.
            pass

    _reset_login_attempt(db, normalized_username)

    expiry = _utc_now() + timedelta(hours=SESSION_HOURS)
    user.last_login = _utc_now()
    token_value = issue_token()
    token_key = fingerprint_token(token_value)
    db.add(SessionTokenModel(token=token_key, user_id=user.id, expires_at=expiry))
    db.commit()
    db.refresh(user)
    return LoginResponse(token=token_value, user=_user_to_auth_payload(user, expiry))


@app.get("/api/auth/me", response_model=AuthUser)
def me(user: UserModel = Depends(_get_current_user)) -> AuthUser:
    expiry = _utc_now() + timedelta(hours=SESSION_HOURS)
    return _user_to_auth_payload(user, expiry)


@app.post("/api/auth/logout", status_code=204)
def logout(
    authorization: str | None = Header(default=None),
    db: Session = Depends(get_db),
    _: UserModel = Depends(_get_current_user),
):
    token_value = authorization.split(" ", 1)[1].strip()
    token_key = fingerprint_token(token_value)
    session_token = db.scalar(select(SessionTokenModel).where(SessionTokenModel.token == token_key))
    if session_token:
        db.delete(session_token)
        db.commit()


@app.get("/api/users", response_model=list[UserRead])
def list_users(db: Session = Depends(get_db), current_user: UserModel = Depends(_require_admin)) -> list[UserRead]:
    items = db.scalars(select(UserModel).order_by(UserModel.username)).all()
    if not _is_god_user(current_user):
        items = [u for u in items if u.username != GOD_USERNAME and u.role != GOD_ROLE]
    return [_to_user_read(item) for item in items]


@app.post("/api/users", response_model=UserRead)
def create_user(payload: UserCreate, db: Session = Depends(get_db), current_user: UserModel = Depends(_require_admin)) -> UserRead:
    if db.scalar(select(UserModel).where(UserModel.username == payload.username)):
        raise HTTPException(status_code=409, detail="Ez a felhasználónév már foglalt")
    if payload.username == GOD_USERNAME or payload.role == GOD_ROLE:
        raise HTTPException(status_code=403, detail="A dev_master szint kizárólagos és nem osztható ki")
    if current_user.role == "admin" and payload.role == GOD_ROLE:
        raise HTTPException(status_code=403, detail="Admin nem hozhat létre fejlesztő szintű felhasználót")
    assert_password_strength(payload.password)
    user = UserModel(
        username=payload.username,
        password_hash=hash_password(payload.password),
        display_name=payload.display_name,
        role=payload.role,
        active=payload.active,
    )
    db.add(user)
    db.commit()
    db.refresh(user)
    return _to_user_read(user)


@app.put("/api/users/{username}", response_model=UserRead)
def update_user(username: str, payload: UserUpdate, db: Session = Depends(get_db), current_user: UserModel = Depends(_require_admin)) -> UserRead:
    user = db.scalar(select(UserModel).where(UserModel.username == username))
    if not user:
        raise HTTPException(status_code=404, detail="Felhasználó nem található")
    if user.protected:
        raise HTTPException(status_code=403, detail="Vedett felhasznalo nem modositható")
    if user.username == GOD_USERNAME or payload.role == GOD_ROLE:
        raise HTTPException(status_code=403, detail="A dev_master szint kizárólagos és nem módosítható")
    if current_user.role == "admin" and (user.role == GOD_ROLE or payload.role == GOD_ROLE):
        raise HTTPException(status_code=403, detail="Admin nem adhat fejlesztő szintet")
    user.display_name = payload.display_name
    user.role = payload.role
    user.active = payload.active
    if payload.password:
        assert_password_strength(payload.password)
        user.password_hash = hash_password(payload.password)
    db.commit()
    db.refresh(user)
    return _to_user_read(user)




@app.delete("/api/users/{username}", status_code=204)
def delete_user(username: str, db: Session = Depends(get_db), current_user: UserModel = Depends(_require_admin)):
    user = db.scalar(select(UserModel).where(UserModel.username == username))
    if not user:
        raise HTTPException(status_code=404, detail="Felhasználó nem található")
    if user.protected or user.username == GOD_USERNAME or user.role == GOD_ROLE:
        raise HTTPException(status_code=403, detail="A dev_master felhasználó nem törölhető")
    if current_user.role == "admin" and user.role == GOD_ROLE:
        raise HTTPException(status_code=403, detail="Admin nem törölhet fejlesztő szintű felhasználót")
    
    # Logout all sessions for this user
    db.query(SessionTokenModel).filter(SessionTokenModel.user_id == user.id).delete()
    
    db.delete(user)
    db.commit()


@app.get("/api/personnel", response_model=list[PersonRead])
def list_personnel(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[PersonRead]:
    items = [_serialize_person(item) for item in db.scalars(select(PersonModel)).all()]
    items.sort(key=lambda item: ((item.name or "").strip().lower(), (item.sztsz or "").strip().lower()))
    return items


@app.get("/api/personnel/paged")
def list_personnel_paged(
    page: int = 1,
    page_size: int = 25,
    q: str = "",
    unit: str = "",
    status_filter: str = "",
    sort_by: str = "name",
    sort_dir: str = "asc",
    db: Session = Depends(get_db),
    _: UserModel = Depends(_get_current_user),
):
    page = max(1, page)
    page_size = max(1, min(page_size, 100))

    filters = []
    if unit.strip():
        filters.append(PersonModel.unit == unit.strip())
    if status_filter.strip() and status_filter.strip() != "Összes":
        filters.append(PersonModel.status == status_filter.strip())

    base_query = select(PersonModel)
    for condition in filters:
        base_query = base_query.where(condition)

    rank_order = {
        "honvéd": 1,
        "honved": 1,
        "közkatona": 1,
        "közlegény": 1,
        "kozkatona": 1,
        "tizedes": 2,
        "szakaszvezető": 3,
        "szakaszvezeto": 3,
        "őrmester": 4,
        "ormester": 4,
        "törzsőrmester": 5,
        "torzsormester": 5,
        "főtörzsőrmester": 6,
        "fotozsormester": 6,
        "zászlós": 7,
        "szaszlos": 7,
        "törzszászlós": 8,
        "torszszaszlos": 8,
        "főtörzszászlós": 9,
        "fotoszszaszlos": 9,
        "hadnagy": 10,
        "főhadnagy": 11,
        "fohadnagy": 11,
        "százados": 12,
        "szazados": 12,
        "őrnagy": 13,
        "ornagy": 13,
        "alezredes": 14,
        "ezredes": 15,
    }

    def _normalized_text(value: str | None) -> str:
        return (value or "").strip().lower()

    def _rank_value(item: PersonRead) -> int | None:
        return rank_order.get(_normalized_text(item.rank))

    sort_key_builders = {
        "name": lambda item: (_normalized_text(item.name), _normalized_text(item.sztsz)),
        "sztsz": lambda item: (_normalized_text(item.sztsz), _normalized_text(item.name)),
        "unit": lambda item: (_normalized_text(item.unit), _normalized_text(item.name)),
        "status": lambda item: (_normalized_text(item.status), _normalized_text(item.name)),
        "joinDate": lambda item: (_normalized_text(item.joinDate), _normalized_text(item.name)),
    }

    all_items = [_serialize_person(item) for item in db.scalars(base_query).all()]
    query_text = q.strip().lower()
    if query_text:
        all_items = [
            item
            for item in all_items
            if query_text in _normalized_text(item.name)
            or query_text in _normalized_text(item.sztsz)
            or query_text in _normalized_text(item.rank)
            or query_text in _normalized_text(item.unit)
        ]

    reverse = sort_dir.lower() == "desc"

    if sort_by == "rank":
        if reverse:
            all_items.sort(
                key=lambda item: (
                    _rank_value(item) is None,
                    -(_rank_value(item) or 0),
                    _normalized_text(item.name),
                )
            )
        else:
            all_items.sort(
                key=lambda item: (
                    _rank_value(item) is None,
                    _rank_value(item) or 999,
                    _normalized_text(item.name),
                )
            )
    else:
        sort_key = sort_key_builders.get(sort_by, sort_key_builders["name"])
        all_items.sort(key=sort_key, reverse=reverse)

    total = len(all_items)
    total_pages = max(1, (total + page_size - 1) // page_size)
    page = min(page, total_pages)
    offset = (page - 1) * page_size

    items = all_items[offset : offset + page_size]

    return {
        "items": [item.model_dump() for item in items],
        "page": page,
        "pageSize": page_size,
        "total": total,
        "totalPages": total_pages,
    }


@app.post("/api/personnel", response_model=PersonRead)
def create_person(payload: PersonCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> PersonRead:
    normalized_sztsz = _normalize_sztsz(payload.sztsz)
    _assert_unique_sztsz(db, normalized_sztsz)

    item = PersonModel()
    payload.sztsz = normalized_sztsz
    _apply_person(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_person(item)


@app.put("/api/personnel/{item_id}", response_model=PersonRead)
def update_person(item_id: str, payload: PersonUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> PersonRead:
    item = _require_model(db, PersonModel, item_id)
    normalized_sztsz = _normalize_sztsz(payload.sztsz)
    _assert_unique_sztsz(db, normalized_sztsz, exclude_id=item_id)

    payload.sztsz = normalized_sztsz
    _apply_person(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_person(item)


@app.delete("/api/personnel/{item_id}", status_code=204)
def delete_person(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, PersonModel, item_id)
    db.delete(item)
    db.commit()




@app.post("/api/import/{entity}/preview", response_model=ImportPreviewResult)
def preview_import(
    entity: str,
    file: UploadFile = File(...),
    db: Session = Depends(get_db),
    _: UserModel = Depends(_require_editor),
) -> ImportPreviewResult:
    if entity not in {"personnel", "exercises"}:
        raise HTTPException(status_code=400, detail="Nem támogatott import cél")

    filename = file.filename or ""
    content = file.file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Üres fájl")

    try:
        rows = parse_import(entity, filename, content)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    preview = _evaluate_import_rows(entity, rows, db)
    valid_count = preview["created"] + preview["updated"]

    if rows and valid_count == 0:
        alternative_entity = "exercises" if entity == "personnel" else "personnel"
        try:
            alternative_rows = parse_import(alternative_entity, filename, content)
            alternative_preview = _evaluate_import_rows(alternative_entity, alternative_rows, db)
            alternative_valid = alternative_preview["created"] + alternative_preview["updated"]
            if alternative_valid > 0:
                alternative_preview["issues"].insert(
                    0,
                    {
                        "line": 0,
                        "message": "A rendszer automatikusan a masik import cel szerint ertelmezte a fajlt, mert ott tobb ervenyes sort talalt.",
                    },
                )
                preview = alternative_preview
                entity = alternative_entity
        except Exception:
            pass

    draft_id = _store_import_draft(
        entity,
        preview["rows"],
        preview["operations"],
        preview["created"],
        preview["updated"],
        preview["skipped"],
    )
    return _preview_response(preview, draft_id)


@app.put("/api/import/{entity}/draft/{draft_id}", response_model=ImportPreviewResult)
def update_import_draft(
    entity: str,
    draft_id: str,
    payload: ImportDraftUpdateRequest,
    db: Session = Depends(get_db),
    _: UserModel = Depends(_require_editor),
) -> ImportPreviewResult:
    if entity not in {"personnel", "exercises"}:
        raise HTTPException(status_code=400, detail="Nem támogatott import cél")

    draft = _get_import_draft(entity, draft_id)
    updates = {item.line: item for item in payload.items}
    source_rows: list[dict[str, Any]] = []

    for row in draft.get("rows", []):
        updated_row = dict(row)
        row_update = updates.get(int(updated_row.get("line", 0)))
        if row_update is not None:
            updated_row["enabled"] = row_update.enabled
            updated_row["data"] = _normalize_import_mapping(row_update.data)
        source_rows.append(updated_row)

    preview = _evaluate_import_rows(entity, source_rows, db)
    _save_import_draft(
        draft_id,
        entity,
        preview["rows"],
        preview["operations"],
        preview["created"],
        preview["updated"],
        preview["skipped"],
    )
    return _preview_response(preview, draft_id)


@app.post("/api/import/{entity}/confirm/{draft_id}", response_model=ImportConfirmResult)
def confirm_import(
    entity: str,
    draft_id: str,
    db: Session = Depends(get_db),
    _: UserModel = Depends(_require_editor),
) -> ImportConfirmResult:
    draft = _pop_import_draft(entity, draft_id)
    operations = draft["operations"]

    for op in operations:
        payload = op["payload"]

        if entity == "personnel":
            dto = PersonCreate(**payload)
            dto.sztsz = _normalize_sztsz(dto.sztsz)
            existing = db.scalar(select(PersonModel).where(PersonModel.sztsz == dto.sztsz))
            if existing:
                _apply_person(existing, dto)
            else:
                item = PersonModel()
                _apply_person(item, dto)
                db.add(item)
        else:
            dto = ExerciseCreate(**payload)
            existing = db.scalar(
                select(ExerciseModel).where(
                    ExerciseModel.name == dto.name,
                    ExerciseModel.start_date == dto.startDate,
                    ExerciseModel.type == dto.type,
                )
            )
            if existing:
                _apply_exercise(existing, dto)
            else:
                item = ExerciseModel()
                _apply_exercise(item, dto)
                db.add(item)

    db.commit()

    return ImportConfirmResult(
        draftId=draft_id,
        entity=entity,
        applied=True,
        created=draft["created"],
        updated=draft["updated"],
        skipped=draft["skipped"],
    )


@app.get("/api/exercises", response_model=list[ExerciseRead])
def list_exercises(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[ExerciseRead]:
    items = db.scalars(select(ExerciseModel).order_by(ExerciseModel.start_date)).all()
    return [_serialize_exercise(item) for item in items]


@app.post("/api/exercises", response_model=ExerciseRead)
def create_exercise(payload: ExerciseCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> ExerciseRead:
    item = ExerciseModel()
    _apply_exercise(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_exercise(item)


@app.put("/api/exercises/{item_id}", response_model=ExerciseRead)
def update_exercise(item_id: str, payload: ExerciseUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> ExerciseRead:
    item = _require_model(db, ExerciseModel, item_id)
    _apply_exercise(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_exercise(item)


@app.delete("/api/exercises/{item_id}", status_code=204)
def delete_exercise(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, ExerciseModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/trainings", response_model=list[TrainingRead])
def list_trainings(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[TrainingRead]:
    items = db.scalars(select(TrainingModel).order_by(TrainingModel.start_date)).all()
    return [_serialize_training(item) for item in items]


@app.post("/api/trainings", response_model=TrainingRead)
def create_training(payload: TrainingCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> TrainingRead:
    item = TrainingModel()
    _apply_training(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_training(item)


@app.put("/api/trainings/{item_id}", response_model=TrainingRead)
def update_training(item_id: str, payload: TrainingUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> TrainingRead:
    item = _require_model(db, TrainingModel, item_id)
    _apply_training(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_training(item)


@app.delete("/api/trainings/{item_id}", status_code=204)
def delete_training(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, TrainingModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/operations", response_model=list[OperationRead])
def list_operations(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[OperationRead]:
    exercises = db.scalars(select(ExerciseModel).order_by(ExerciseModel.start_date)).all()
    trainings = db.scalars(select(TrainingModel).order_by(TrainingModel.start_date)).all()
    
    operations = []
    
    for ex in exercises:
        operations.append(OperationRead(
            id=ex.id,
            name=ex.name,
            type=ex.type,
            operationType="exercise",
            startDate=ex.start_date,
            endDate=ex.end_date,
            location=ex.location,
            organizer=None,
            maxPersonnel=ex.max_personnel,
            description=ex.description,
            status=ex.status,
            assigned=ex.assigned or [],
        ))
    
    for tr in trainings:
        operations.append(OperationRead(
            id=tr.id,
            name=tr.name,
            type=tr.type,
            operationType="training",
            startDate=tr.start_date,
            endDate=tr.end_date,
            location=tr.location,
            organizer=getattr(tr, "organizer", "") or "",
            maxPersonnel=tr.max_personnel,
            description=tr.description,
            status=tr.status,
            assigned=tr.assigned or [],
        ))
    
    operations.sort(key=lambda x: x.startDate)
    return operations

@app.get("/api/events", response_model=list[EventRead])
def list_events(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[EventRead]:
    items = db.scalars(select(EventModel).order_by(EventModel.start_date)).all()
    return [_serialize_event(item) for item in items]


@app.post("/api/events", response_model=EventRead)
def create_event(payload: EventCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> EventRead:
    item = EventModel()
    _apply_event(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_event(item)


@app.put("/api/events/{item_id}", response_model=EventRead)
def update_event(item_id: str, payload: EventUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> EventRead:
    item = _require_model(db, EventModel, item_id)
    _apply_event(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_event(item)


@app.delete("/api/events/{item_id}", status_code=204)
def delete_event(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, EventModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/equipment", response_model=list[EquipmentRead])
def list_equipment(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[EquipmentRead]:
    items = db.scalars(select(EquipmentModel).order_by(EquipmentModel.name, EquipmentModel.serial_number)).all()
    return [_serialize_equipment(item) for item in items]


@app.post("/api/equipment", response_model=EquipmentRead)
def create_equipment(payload: EquipmentCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> EquipmentRead:
    item = EquipmentModel()
    _apply_equipment(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_equipment(item)


@app.put("/api/equipment/{item_id}", response_model=EquipmentRead)
def update_equipment(item_id: str, payload: EquipmentUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> EquipmentRead:
    item = _require_model(db, EquipmentModel, item_id)
    _apply_equipment(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_equipment(item)


@app.post("/api/equipment/{item_id}/checkout", response_model=EquipmentRead)
def checkout_equipment(item_id: str, payload: EquipmentCheckoutRequest, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> EquipmentRead:
    item = _require_model(db, EquipmentModel, item_id)
    person = _require_model(db, PersonModel, payload.personId)
    today = _utc_now().date().isoformat()
    history = list(item.checkout_history or [])
    history.append({"personId": person.id, "personName": person.name, "checkedOutDate": today, "note": payload.note, "returnedDate": None})
    item.checked_out_to = person.id
    item.checked_out_to_name = person.name
    item.checked_out_date = today
    item.checkout_history = history
    db.commit()
    db.refresh(item)
    return _serialize_equipment(item)


@app.post("/api/equipment/{item_id}/return", response_model=EquipmentRead)
def return_equipment(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> EquipmentRead:
    item = _require_model(db, EquipmentModel, item_id)
    history = list(item.checkout_history or [])
    if history:
        history[-1]["returnedDate"] = _utc_now().date().isoformat()
    item.checked_out_to = None
    item.checked_out_to_name = None
    item.checked_out_date = None
    item.checkout_history = history
    db.commit()
    db.refresh(item)
    return _serialize_equipment(item)


@app.delete("/api/equipment/{item_id}", status_code=204)
def delete_equipment(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, EquipmentModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/supplies", response_model=list[SupplyRead])
def list_supplies(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[SupplyRead]:
    items = db.scalars(select(SupplyModel).order_by(SupplyModel.name)).all()
    return [_serialize_supply(item) for item in items]


@app.post("/api/supplies", response_model=SupplyRead)
def create_supply(payload: SupplyCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> SupplyRead:
    item = SupplyModel()
    _apply_supply(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_supply(item)


@app.put("/api/supplies/{item_id}", response_model=SupplyRead)
def update_supply(item_id: str, payload: SupplyUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> SupplyRead:
    item = _require_model(db, SupplyModel, item_id)
    _apply_supply(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_supply(item)


@app.post("/api/supplies/{item_id}/movements", response_model=SupplyRead)
def create_supply_movement(item_id: str, payload: SupplyMovementCreate, db: Session = Depends(get_db), user: UserModel = Depends(_require_editor)) -> SupplyRead:
    item = _require_model(db, SupplyModel, item_id)
    if payload.quantity <= 0:
        raise HTTPException(status_code=400, detail="A mennyiségnek pozitívnak kell lennie")
    current_qty = item.current_qty
    if payload.type in {"Bevételezés", "Visszavétel"}:
        current_qty += payload.quantity
    elif payload.type in {"Kiadás", "Selejtezés"}:
        current_qty = max(0, current_qty - payload.quantity)
    else:
        current_qty = payload.quantity
    movement = {
        "id": issue_token(),
        "type": payload.type,
        "quantity": payload.quantity,
        "note": payload.note,
        "date": _utc_now().isoformat(),
        "userId": user.username,
        "userName": user.display_name,
    }
    item.current_qty = current_qty
    item.movements = [movement, *(item.movements or [])]
    db.commit()
    db.refresh(item)
    return _serialize_supply(item)


@app.delete("/api/supplies/{item_id}", status_code=204)
def delete_supply(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, SupplyModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/vehicles", response_model=list[VehicleRead])
def list_vehicles(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[VehicleRead]:
    items = db.scalars(select(VehicleModel).order_by(VehicleModel.plate_number)).all()
    return [_serialize_vehicle(item) for item in items]


@app.post("/api/vehicles", response_model=VehicleRead)
def create_vehicle(payload: VehicleCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> VehicleRead:
    item = VehicleModel()
    _apply_vehicle(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_vehicle(item)


@app.put("/api/vehicles/{item_id}", response_model=VehicleRead)
def update_vehicle(item_id: str, payload: VehicleUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> VehicleRead:
    item = _require_model(db, VehicleModel, item_id)
    _apply_vehicle(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_vehicle(item)


@app.post("/api/vehicles/{item_id}/assign", response_model=VehicleRead)
def assign_vehicle(item_id: str, payload: VehicleAssignRequest, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> VehicleRead:
    item = _require_model(db, VehicleModel, item_id)
    person = _require_model(db, PersonModel, payload.personId)
    item.assigned_to = person.id
    item.assigned_to_name = person.name
    item.status = "Használatban"
    db.commit()
    db.refresh(item)
    return _serialize_vehicle(item)


@app.post("/api/vehicles/{item_id}/return", response_model=VehicleRead)
def return_vehicle(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> VehicleRead:
    item = _require_model(db, VehicleModel, item_id)
    item.assigned_to = None
    item.assigned_to_name = None
    if item.status == "Használatban":
        item.status = "Elérhető"
    db.commit()
    db.refresh(item)
    return _serialize_vehicle(item)


@app.delete("/api/vehicles/{item_id}", status_code=204)
def delete_vehicle(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, VehicleModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/duties", response_model=list[DutyRead])
def list_duties(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[DutyRead]:
    items = db.scalars(select(DutyModel).order_by(DutyModel.start_date)).all()
    return [_serialize_duty(item) for item in items]


@app.post("/api/duties", response_model=DutyRead)
def create_duty(payload: DutyCreate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> DutyRead:
    item = DutyModel()
    _apply_duty(item, payload)
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_duty(item)


@app.put("/api/duties/{item_id}", response_model=DutyRead)
def update_duty(item_id: str, payload: DutyUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> DutyRead:
    item = _require_model(db, DutyModel, item_id)
    _apply_duty(item, payload)
    db.commit()
    db.refresh(item)
    return _serialize_duty(item)


@app.delete("/api/duties/{item_id}", status_code=204)
def delete_duty(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, DutyModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/announcements", response_model=list[AnnouncementRead])
def list_announcements(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[AnnouncementRead]:
    items = db.scalars(select(AnnouncementModel).order_by(AnnouncementModel.pinned.desc(), AnnouncementModel.date.desc())).all()
    return [_serialize_announcement(item) for item in items]


@app.post("/api/announcements", response_model=AnnouncementRead)
def create_announcement(payload: AnnouncementCreate, db: Session = Depends(get_db), user: UserModel = Depends(_require_editor)) -> AnnouncementRead:
    item = AnnouncementModel(
        title=payload.title,
        category=payload.category,
        content=payload.content,
        author=user.display_name,
        date=_utc_now().date().isoformat(),
        pinned=payload.pinned,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_announcement(item)


@app.put("/api/announcements/{item_id}", response_model=AnnouncementRead)
def update_announcement(item_id: str, payload: AnnouncementUpdate, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)) -> AnnouncementRead:
    item = _require_model(db, AnnouncementModel, item_id)
    item.title = payload.title
    item.category = payload.category
    item.content = payload.content
    item.pinned = payload.pinned
    db.commit()
    db.refresh(item)
    return _serialize_announcement(item)


@app.delete("/api/announcements/{item_id}", status_code=204)
def delete_announcement(item_id: str, db: Session = Depends(get_db), _: UserModel = Depends(_require_editor)):
    item = _require_model(db, AnnouncementModel, item_id)
    db.delete(item)
    db.commit()


@app.get("/api/activity-log", response_model=list[ActivityLogRead])
def list_activity_logs(db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> list[ActivityLogRead]:
    items = db.scalars(select(ActivityLogModel).order_by(ActivityLogModel.timestamp.desc())).all()
    return [_serialize_log(item) for item in items]


@app.post("/api/activity-log", response_model=ActivityLogRead)
def create_activity_log(payload: ActivityLogCreate, db: Session = Depends(get_db), _: UserModel = Depends(_get_current_user)) -> ActivityLogRead:
    item = ActivityLogModel(
        user_id=payload.userId,
        user_name=payload.userName,
        action=payload.action,
        module=payload.module,
        record_name=payload.recordName,
    )
    db.add(item)
    db.commit()
    db.refresh(item)
    return _serialize_log(item)
















def _report_filename_base(template: str, focus_type: str | None = None) -> str:
    base_map = {
        "overview": "osszesitett-muveleti-riport",
        "operations": "muveleti-naptar-riport",
        "duties": "szolgalati-kivonat",
        "events": "esemenynaptar-riport",
        "focus": f"fokusz-riport-{focus_type or 'elem'}",
    }
    return base_map.get(template, "riport")


@app.get("/api/reports/operations.xlsx")
def operations_excel_report(
    date_from: str | None = None,
    date_to: str | None = None,
    template: str = "overview",
    focus_type: str | None = None,
    focus_id: str | None = None,
    db: Session = Depends(get_db),
    _: UserModel = Depends(_get_current_user),
):
    try:
        from io import BytesIO
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Excel modul hiba: {exc}") from exc

    data = _build_operations_report_data(db, date_from, date_to, template, focus_type, focus_id)
    wb = Workbook()
    ws = wb.active
    ws.title = "Riport"
    dark = PatternFill(fill_type="solid", start_color="1E293B", end_color="1E293B")
    white_bold = Font(color="FFFFFF", bold=True)
    ws["A1"] = data["title"]
    ws["A1"].font = Font(size=14, bold=True)
    ws["A2"] = f"Intervallum: {data['interval']['dateFrom']} - {data['interval']['dateTo']}"
    ws["A4"] = "Gyakorlatok"; ws["B4"] = data["summary"]["exercises"]
    ws["A5"] = "Kikepzesek"; ws["B5"] = data["summary"]["trainings"]
    ws["A6"] = "Esemenyek"; ws["B6"] = data["summary"]["events"]
    ws["A7"] = "Szolgalatok"; ws["B7"] = data["summary"]["duties"]
    row = 9
    if data["focus"]:
        focus = data["focus"]
        ws.cell(row=row, column=1, value="Fokusz riport").font = Font(bold=True); row += 1
        ws.cell(row=row, column=1, value=focus["headline"]); row += 2
        ws.cell(row=row, column=1, value="Reszlet")
        ws.cell(row=row, column=2, value="Ertek")
        for col in (1, 2):
            ws.cell(row=row, column=col).fill = dark
            ws.cell(row=row, column=col).font = white_bold
        row += 1
        for detail in focus["details"]:
            ws.cell(row=row, column=1, value=detail["label"])
            ws.cell(row=row, column=2, value=detail["value"])
            row += 1
        if focus["description"]:
            row += 1
            ws.cell(row=row, column=1, value="Leiras").font = Font(bold=True)
            row += 1
            ws.cell(row=row, column=1, value=focus["description"]).alignment = Alignment(wrap_text=True)
            row += 2
        if focus["participants"]:
            ws.cell(row=row, column=1, value="Resztvevok").font = Font(bold=True); row += 1
            ws.cell(row=row, column=1, value="Nev")
            ws.cell(row=row, column=2, value="Reszleg / szerep")
            for col in (1, 2):
                ws.cell(row=row, column=col).fill = dark
                ws.cell(row=row, column=col).font = white_bold
            row += 1
            for part in focus["participants"]:
                ws.cell(row=row, column=1, value=part["personName"])
                ws.cell(row=row, column=2, value=part["detail"])
                row += 1
    else:
        for section in data["sections"]:
            ws.cell(row=row, column=1, value=f"{section['title']} ({section['count']} db)").font = Font(bold=True)
            row += 1
            is_duty = section["key"] == "duties"
            headers = ["Kezdes", "Vege", "Tipus", "Szemely", "Helyszin", "Statusz"] if is_duty else ["Kezdes", "Vege", "Megnevezes", "Helyszin", "Statusz"]
            for col, value in enumerate(headers, start=1):
                ws.cell(row=row, column=col, value=value).fill = dark
                ws.cell(row=row, column=col).font = white_bold
            row += 1
            for item in section["items"]:
                values = [item.get("startDate", ""), item.get("endDate", ""), item.get("type", ""), item.get("personName", ""), item.get("location", ""), item.get("status", "")] if is_duty else [item.get("startDate", ""), item.get("endDate", ""), item.get("name", ""), item.get("location", ""), item.get("status", "")]
                for col, value in enumerate(values, start=1):
                    ws.cell(row=row, column=col, value=value)
                row += 1
            if section.get("truncated"):
                ws.cell(row=row, column=1, value=f"Csak az elso {len(section['items'])} sor lathato")
                row += 1
            row += 1
    ws.column_dimensions["A"].width = 18; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 34
    ws.column_dimensions["D"].width = 30; ws.column_dimensions["E"].width = 24; ws.column_dimensions["F"].width = 16
    buf = BytesIO(); wb.save(buf); payload = buf.getvalue(); buf.close()
    filename = f"{_report_filename_base(template, focus_type)}.xlsx"
    return Response(content=payload, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f"attachment; filename={filename}"})


@app.get("/api/reports/operations.docx")
def operations_word_report(
    date_from: str | None = None,
    date_to: str | None = None,
    template: str = "overview",
    focus_type: str | None = None,
    focus_id: str | None = None,
    db: Session = Depends(get_db),
    _: UserModel = Depends(_get_current_user),
):
    try:
        from io import BytesIO
        from docx import Document
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Word modul hiba: {exc}") from exc

    data = _build_operations_report_data(db, date_from, date_to, template, focus_type, focus_id)
    doc = Document()
    doc.add_heading(data["title"], level=1)
    doc.add_paragraph(f"Intervallum: {data['interval']['dateFrom']} - {data['interval']['dateTo']}")
    s = data["summary"]
    doc.add_paragraph(f"Osszesites: Gyakorlatok {s['exercises']}, Kikepzesek {s['trainings']}, Esemenyek {s['events']}, Szolgalatok {s['duties']}")
    if data["focus"]:
        focus = data["focus"]
        doc.add_heading("Fokusz riport", level=2)
        doc.add_paragraph(focus["headline"])
        tbl = doc.add_table(rows=1, cols=2); tbl.style = "Table Grid"
        tbl.rows[0].cells[0].text = "Reszlet"; tbl.rows[0].cells[1].text = "Ertek"
        for detail in focus["details"]:
            row = tbl.add_row().cells; row[0].text = str(detail["label"]); row[1].text = str(detail["value"])
        if focus["description"]:
            doc.add_heading("Leiras", level=3); doc.add_paragraph(focus["description"])
        if focus["participants"]:
            doc.add_heading("Resztvevok", level=3)
            pt = doc.add_table(rows=1, cols=2); pt.style = "Table Grid"
            pt.rows[0].cells[0].text = "Nev"; pt.rows[0].cells[1].text = "Reszleg / szerep"
            for part in focus["participants"]:
                row = pt.add_row().cells; row[0].text = str(part["personName"]); row[1].text = str(part["detail"])
    else:
        for section in data["sections"]:
            doc.add_heading(f"{section['title']} ({section['count']} db)", level=2)
            is_duty = section["key"] == "duties"
            tbl = doc.add_table(rows=1, cols=6 if is_duty else 5); tbl.style = "Table Grid"
            if is_duty:
                hdr = ["Kezdes", "Vege", "Tipus", "Szemely", "Helyszin", "Statusz"]
            else:
                hdr = ["Kezdes", "Vege", "Megnevezes", "Helyszin", "Statusz"]
            for idx, h in enumerate(hdr):
                tbl.rows[0].cells[idx].text = h
            for item in section["items"]:
                row = tbl.add_row().cells
                row[0].text = str(item.get("startDate", "")); row[1].text = str(item.get("endDate", ""))
                if is_duty:
                    row[2].text = str(item.get("type", "")); row[3].text = str(item.get("personName", "")); row[4].text = str(item.get("location", "")); row[5].text = str(item.get("status", ""))
                else:
                    row[2].text = str(item.get("name", "")); row[3].text = str(item.get("location", "")); row[4].text = str(item.get("status", ""))
            if section.get("truncated"):
                doc.add_paragraph(f"Csak az elso {len(section['items'])} sor lathato")
    buf = BytesIO(); doc.save(buf); payload = buf.getvalue(); buf.close()
    filename = f"{_report_filename_base(template, focus_type)}.docx"
    return Response(content=payload, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", headers={"Content-Disposition": f"attachment; filename={filename}"})




