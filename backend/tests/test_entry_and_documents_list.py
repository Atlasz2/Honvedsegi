"""2. betekintés (Várpalota): típusos okmányok, lejárati lista, beléptetés/átjelentő, jegyzőkönyv."""
from __future__ import annotations

import io
from datetime import date, timedelta

from openpyxl import load_workbook

CARD = "Szolgálati igazolvány"


def _person(client, headers, name, sztsz, rank="honvéd"):
    r = client.post("/api/personnel", json={"name": name, "sztsz": sztsz, "rank": rank, "unit": "1. század", "status": "Aktív"}, headers=headers)
    assert r.status_code == 200, r.text
    return r.json()["id"]


def _sheet(response):
    assert response.status_code == 200, response.text
    return [list(r) for r in load_workbook(io.BytesIO(response.content)).active.iter_rows(values_only=True)]


def test_rank_mismatch_and_filtered_list(client, admin_headers):
    pid = _person(client, admin_headers, "Kártyás Kata", "18700001", rank="őrvezető")
    soon = (date.today() + timedelta(days=10)).isoformat()
    client.post(f"/api/documents/personnel/{pid}", json={"name": CARD, "identifier": "SZ-1", "expiryDate": soon, "rankOnCard": "honvéd"}, headers=admin_headers)
    client.post(f"/api/documents/personnel/{pid}", json={"name": "Vezetői engedély (sárga kártya)", "expiryDate": "2099-01-01"}, headers=admin_headers)

    all_rows = client.get("/api/documents", headers=admin_headers).json()
    mine = [r for r in all_rows if r["personnelId"] == pid]
    assert len(mine) == 2
    assert client.get("/api/documents", params={"name": CARD}, headers=admin_headers).json()[0]["name"] == "Kártyás Kata"
    mismatch = client.get("/api/documents", params={"mismatch": True}, headers=admin_headers).json()
    assert [r["personnelId"] for r in mismatch] == [pid] and mismatch[0]["rankMismatch"]
    assert all(r["daysUntilExpiry"] <= 30 for r in client.get("/api/documents", params={"days": 30}, headers=admin_headers).json())

    export = _sheet(client.get("/api/documents/export.xlsx", params={"mismatch": True}, headers=admin_headers))
    assert any("ÚJ KÁRTYA KELL" in str(cell) for row in export for cell in row)


def test_entry_meals_and_deadline(client, admin_headers):
    pid = _person(client, admin_headers, "Belépő Béla", "18700002")
    _person(client, admin_headers, "Lemondó Lajos", "18700003")
    client.post(f"/api/documents/personnel/{pid}", json={"name": CARD, "identifier": "SZ-77", "expiryDate": "2099-06-01"}, headers=admin_headers)
    start = (date.today() + timedelta(days=6)).isoformat()
    end = (date.today() + timedelta(days=7)).isoformat()
    ex = client.post("/api/exercises", json={"name": "Beléptetős gyakorlat", "type": "Gyakorlat", "startDate": start, "endDate": end,
                                              "location": "Várpalota", "maxPersonnel": 5, "description": "", "status": "Tervezett",
                                              "assigned": [{"personId": pid, "personName": "Belépő Béla"}]}, headers=admin_headers).json()

    rows = _sheet(client.get(f"/api/entry/{ex['id']}/entry.xlsx", headers=admin_headers))
    assert any("SZ-77" in row for row in rows)

    meals = _sheet(client.get(f"/api/entry/{ex['id']}/meals.xlsx", headers=admin_headers))
    assert meals[-1][0] == "Összesen" and meals[-1][3:5] == [1, 1]

    deadlines = client.get("/api/entry/deadlines", headers=admin_headers).json()
    mine = next(d for d in deadlines if d["exerciseId"] == ex["id"])
    assert mine["daysLeft"] == 1 and mine["participants"] == 1  # kezdés - 5 nap = holnap


def test_discharge_protocol_docx(client, admin_headers):
    pid = _person(client, admin_headers, "Leszerelő Laci", "18700004")
    eq = client.post("/api/equipment", json={"name": "Sisak", "category": "Védőfelszerelés", "serialNumber": "S-1", "condition": "Jó"}, headers=admin_headers).json()
    client.post(f"/api/equipment/{eq['id']}/checkout", json={"personId": pid}, headers=admin_headers)

    r = client.post(f"/api/entry/protocol/{pid}", json={"items": [{"equipmentId": eq["id"], "returned": False, "amount": 12000}]}, headers=admin_headers)
    assert r.status_code == 200 and r.content[:2] == b"PK"  # DOCX = zip
    other = _person(client, admin_headers, "Más Miska", "18700005")
    assert client.post(f"/api/entry/protocol/{other}", json={"items": [{"equipmentId": eq["id"]}]}, headers=admin_headers).status_code == 400
